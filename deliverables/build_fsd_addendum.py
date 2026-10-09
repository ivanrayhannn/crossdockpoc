from pathlib import Path
import math
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("FSD_Crossdock_Batch_Uploads.xlsx")
NAVY, BLUE, GRAY, WHITE, TEXT = "1F3A5F", "DCE6F2", "F2F2F2", "FFFFFF", "222222"
BASE, PROP, OPEN, LINE = "E2F0D9", "FFF2CC", "FCE4D6", "B7C9D6"
thin = Side(style="thin", color=LINE)

def setup(ws, widths, portrait=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait" if portrait else "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.45
    ws.oddFooter.center.text = "FSD - Crossdock Master Setting | Draft for review"
    ws.oddFooter.right.text = "Page &[Page] of &[Pages]"
    ws.freeze_panes = "A4"

def row_height(values, widths):
    n = 1
    for i, val in enumerate(values):
        if val is None: continue
        width = widths[min(i, len(widths)-1)]
        n = max(n, math.ceil(len(str(val)) / max(8, width * 1.05)), str(val).count("\n") + 1)
    return min(120, max(24, n * 15 + 8))

def body(cell, val, status=False):
    cell.value = val
    cell.font = Font(name="Calibri", size=10, color=TEXT)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if status:
        s = str(val).lower()
        color = BASE if "baseline" in s or "confirmed" in s else PROP if "proposed" in s or "draft" in s else OPEN
        cell.fill = PatternFill("solid", fgColor=color)

def title(ws, name, subtitle, ncols):
    end = get_column_letter(ncols)
    ws.merge_cells(f"A1:{end}1")
    c = ws["A1"]; c.value = name
    c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{end}2")
    c = ws["A2"]; c.value = subtitle
    c.font = Font(name="Calibri", size=10, italic=True, color="404040")
    c.alignment = Alignment(vertical="top", wrap_text=True); ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 8

def section(ws, row, heading, headers, rows, widths):
    n = len(headers); ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    c = ws.cell(row, 1, heading); c.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 23; row += 1
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col, h); c.font = Font(name="Calibri", size=10, bold=True, color=TEXT)
        c.fill = PatternFill("solid", fgColor=BLUE); c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[row].height = 30; row += 1
    for vals in rows:
        for col, h in enumerate(headers, 1):
            v = vals[col-1] if col-1 < len(vals) else ""
            body(ws.cell(row, col), v, h.lower() in ("status", "requirement status", "decision status"))
        ws.row_dimensions[row].height = row_height(vals, widths); row += 1
    return row + 1

def sheet(name, heading, subtitle, widths, sections):
    ws = wb.create_sheet(name); setup(ws, widths); title(ws, heading, subtitle, len(widths))
    r = 4
    for sh, headers, rows in sections: r = section(ws, r, sh, headers, rows, widths)
    ws.print_title_rows = "1:3"
    return ws

wb = Workbook(); cover = wb.active; cover.title = "Cover"
setup(cover, [27, 74, 28, 24], portrait=True)
cover.merge_cells("A2:D2"); cover["A2"] = "FUNCTIONAL SPECIFICATION DOCUMENT"
cover["A2"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
cover["A3"] = "Crossdock Master Setting - Batch Uploads"
cover["A3"].font = Font(name="Calibri", size=17, bold=True, color=TEXT)
cover["A4"] = "New Service Part System (NSP) | Master > Crossdock Master Setting"
cover["A4"].font = Font(name="Calibri", size=11, color="555555")
meta = [
    ("Document", "FSD - Crossdock Batch Uploads"),
    ("Version", "0.1 (Draft for review)"),
    ("Date", "2026-10-08"),
    ("Scope", "Part List Upload and Mapping per Destination Upload"),
    ("Basis", "FSD_Crossdock_Master_Setting.xlsx v0.2; UISS_Batch_PartListUpload.xlsx v0.2; UISS_Batch_MappingUpload.xlsx v0.2"),
    ("Language", "English; field and button labels follow the existing application specification."),
]
for r, (label, value) in enumerate(meta, 6):
    a, b = cover.cell(r, 1, label), cover.cell(r, 2, value)
    a.font = Font(name="Calibri", size=10, bold=True); a.fill = PatternFill("solid", fgColor=GRAY)
    b.font = Font(name="Calibri", size=10); b.alignment = Alignment(wrap_text=True, vertical="top")
    for c in (a, b): c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cover.row_dimensions[r].height = 32
cover.merge_cells("A14:D14"); cover["A14"] = "Authors and approval (fill in manually)"
cover["A14"].font = Font(name="Calibri", size=11, bold=True, color=WHITE); cover["A14"].fill = PatternFill("solid", fgColor=NAVY)
for r, label in [(15, "Prepared by"), (16, "Reviewed by"), (17, "Approved by")]:
    cover.cell(r, 1, label).font = Font(bold=True); cover.cell(r, 1).fill = PatternFill("solid", fgColor=GRAY)
    cover.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for col in range(1, 5): cover.cell(r, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
cover.merge_cells("A19:D19"); cover["A19"] = "Revision history"
cover["A19"].font = Font(name="Calibri", size=11, bold=True, color=WHITE); cover["A19"].fill = PatternFill("solid", fgColor=NAVY)
for col, v in enumerate(["Version", "Change", "Date", "Owner"], 1):
    c = cover.cell(20, col, v); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor=BLUE)
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
for col, v in enumerate(["0.1", "Initial draft for Part List and Mapping batch uploads; proposals and open decisions are identified.", "2026-10-08", "Ivan"], 1):
    c = cover.cell(21, col, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
cover.row_dimensions[21].height = 45
cover.merge_cells("A23:D23"); cover["A23"] = "Contents - select a section name to open its sheet"
cover["A23"].font = Font(name="Calibri", size=11, bold=True, color=WHITE); cover["A23"].fill = PatternFill("solid", fgColor=NAVY)
toc = [
    ("1 Overview", "1 Overview", "Purpose, scope, actors, terms and source requirements"),
    ("2 Shared Batch Requirements", "2 Shared Requirements", "Common processing, file checks, transaction and status rules"),
    ("3 Part List Upload", "3 Part List Upload", "Template, row validation, save effects and acceptance criteria"),
    ("4 Mapping Upload", "4 Mapping Upload", "Proposed Mapping template, group rules, validation and save behavior"),
    ("5 Data, Results and Audit", "5 Data and Results", "Logical data, batch result and audit requirements"),
    ("6 Messages and Acceptance Criteria", "6 Messages and AC", "Draft user messages and observable acceptance criteria"),
    ("7 Open Points", "7 Open Points", "Decisions needed before implementation"),
]
for r, (label, target, desc) in enumerate(toc, 24):
    c = cover.cell(r, 1, label); c.hyperlink = f"#'{target}'!A1"; c.font = Font(color="0563C1", underline="single")
    cover.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cover.cell(r, 2, desc).alignment = Alignment(wrap_text=True)
    for col in range(1, 5): cover.cell(r, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cover.row_dimensions[r].height = 25
cover.freeze_panes = "A6"; cover.print_area = "A1:D30"

sheet("1 Overview", "1. Overview", "Purpose, scope, actors, source requirements and status of this FSD addendum.", [25,45,72,28], [
("Document scope", ["Topic","Description","Status","Reference"], [
("Purpose","Specify functional behavior for batch upload of Part List data and Mapping per Destination data.","Draft","Current FSD v0.2; UISS Batch v0.2"),
("Included","Upload receipt, file and row/group validation, persistence behavior, result reporting, audit and error handling.","Draft","This addendum"),
("Part List baseline","Template has six columns; .xlsx, 5 MB maximum and 5,000 rows maximum. Execution validation and result behavior are not defined by the existing FSD.","Baseline plus open behavior","Base FSD sections 9 and 10, OP-02"),
("Mapping baseline","Mapping Export and screen rules are specified. Mapping upload UI, file template and processing are not defined by the base FSD.","Proposed function","Base FSD sections 5, 9 and 10"),
("Out of scope","Screen redesign beyond the required upload action; MAD calculation; Sync integration; role provisioning; final physical database schema.","Out of scope","Base FSD v0.2"),
("Status legend","Baseline = stated in existing FSD. Proposed = derived from UISS batch drafts and needs confirmation. Open = cannot be finalized from available sources.","Legend","Apply this status throughout this document."),
]),
("Actors and responsibilities", ["Actor","Function","Access / responsibility","Status"], [
("Procurement user","Part List Upload","Uploads part settings; Add new parts or Edit existing settings for the ongoing period.","Role baseline; processing partly open"),
("D/S user","Mapping Upload","Uploads destination allocations for Active parts in the ongoing period.","Role baseline; upload action proposed"),
("Batch service","Both","Validates workbook, applies approved changes, records processing status and returns a result.","Proposed system behavior"),
("System owner","Both","Maintains batch configuration, storage and logs; handles technical failures.","Operational owner TBD"),
]),
("Terms", ["Term","Meaning","Status","Reference"], [
("UPLOAD_ID","Unique identifier assigned to an uploaded workbook and used to trace one batch invocation.","Proposed","UISS Batch drafts"),
("Ongoing period","Active calculation period. This draft targets only that period; period is not supplied in either template.","Baseline screen rule; batch detail proposed","Base FSD BR-13"),
("Changed group","All uploaded Mapping rows for one Part No form the proposed replacement set for that part's destinations in the ongoing period.","Proposed","UISS Mapping draft"),
("All-or-nothing","If any file/data validation error occurs, no master data is written for the workbook.","Proposed","UISS Batch drafts"),
]),
("High-level process", ["Step","Actor / system","Activity","Outcome"], [
("1","User","Selects a workbook from the relevant upload action and submits it.","File stored and batch request created."),
("2","Batch service","Checks file structure and loads rows for validation.","File errors stop processing; otherwise data is checked."),
("3","Batch service","Validates every row; Mapping rows are also grouped by Part No.","Errors are collected with row/field context."),
("4","Batch service","If data is valid, writes changes using the agreed transaction behavior.","Changed rows/groups saved; identical data skipped."),
("5","Batch service","Creates a result report and finalizes the upload record.","User retrieves status/report and corrects errors if needed."),
]),
])

sheet("2 Shared Requirements", "2. Shared Batch Requirements", "Requirements shared by both upload functions. Proposed rules need owner confirmation.", [16,34,74,28,30], [
("Batch lifecycle", ["ID","Requirement","Description","Status","Reference"], [
("SH-01","Role-based entry","Part List upload is for Procurement users. Mapping upload is intended for D/S users; its action/location must be added to Mapping per Destination.","Part List baseline; Mapping action proposed","Base FSD BR-01 / OP-11"),
("SH-02","Upload receipt","Store the original workbook, create an upload record, assign UPLOAD_ID and start or queue one batch invocation.","Proposed","UISS Batch drafts"),
("SH-03","Batch lifecycle","Each upload moves through RECEIVED, PROCESSING, then COMPLETED or REJECTED. COMPLETED means writes committed or no changes were needed.","Proposed","UISS Batch drafts"),
("SH-04","File checks","Check extension, size, row count, required sheet and headers before master data is written. Function-specific controls are in sections 3 and 4.","Mixed; see sections 3-4","Base FSD section 9; UISS drafts"),
("SH-05","Collect errors","For structurally valid files, validate every row/group and return errors with Excel row, field/group and reason.","Proposed","UISS Batch drafts"),
("SH-06","Atomic save","Draft: any data validation error rejects the workbook and writes no master data. Database failures roll back all writes. Confirm before implementation.","Proposed","UISS Batch drafts"),
("SH-07","Unchanged data","Skip rows/groups whose business values match current master data; do not change audit or Sync Status for skipped data.","Proposed","UISS Batch drafts"),
("SH-08","Period selection","Apply both functions to the ongoing calculation period, inferred by the system and absent from the workbook.","Baseline screen rule; batch implementation proposed","Base FSD BR-13"),
("SH-09","User and time","For changed records, use the authenticated uploader and server timestamp in audit fields.","Proposed","Base FSD OP-10; UISS drafts"),
("SH-10","Result report","Draft result is an .xlsx workbook with Summary, Result, Re-upload and Changes sheets; exact fields and delivery are TBD.","Proposed","UISS Batch drafts"),
("SH-11","Batch result","Receiving UI shows final status and provides report download. Queued/long-running behavior is TBD.","Open","Not defined in base FSD"),
("SH-12","Failure and retry","A rejected upload is not successful. User corrects data and submits a new upload; retry of the same UPLOAD_ID is unspecified.","Proposed / open","UISS Batch drafts"),
]),
("Common file checks", ["Check","Draft rule","Part List status","Mapping status","Notes"], [
("Format","Excel workbook with .xlsx extension.","Baseline","Proposed","Current Mapping export is .xls HTML, not an upload template."),
("Maximum size","5 MB or less.","Baseline","Proposed","Mapping limit is absent from the base FSD."),
("Maximum rows","No more than 5,000 data rows.","Baseline","Proposed","Whether the header counts is TBD."),
("Empty workbook","Reject if there are no data rows.","Open","Open","Base FSD does not state empty-file behavior."),
("Worksheet","Use worksheet named Upload; ignore other sheets.","Proposed","Proposed","Not specified in base FSD."),
("Headers","Require exact agreed template fields and order.","Fields baseline; strict order proposed","Proposed","Decide extra columns, duplicate headers and hidden rows."),
]),
("Batch status definitions", ["Status","Meaning","Status of requirement","Notes"], [
("RECEIVED","File stored and request accepted; batch has not started.","Proposed","Align name with implementation conventions."),
("PROCESSING","Batch has started file and row/group validation.","Proposed","Include start timestamp."),
("COMPLETED","All approved writes committed, or all rows/groups were unchanged.","Proposed","Include row/group counters."),
("REJECTED","File, validation or persistence failed; master changes are not committed under draft atomic-save behavior.","Proposed","Include validation details where possible."),
]),
])

sheet("3 Part List Upload", "3. Part List Upload", "Functional specification for Upload Excel on the Part List screen.", [18,30,67,23,26], [
("Entry and template", ["ID / Field","Requirement","Description","Status","Reference"], [
("PL-01 Entry","Entry point","User opens Part List > Upload Excel, selects a workbook and submits it. Existing modal is only a guide; processing is not defined in the prototype.","Entry baseline; processing open","Base FSD M11 / OP-02"),
("PL-02 Sheet","Worksheet","Draft template worksheet name is Upload.","Proposed","UISS Part List"),
("flag_add_edit","Add/Edit marker","Required for each row. UISS draft uses A for Add and E for Edit; base FSD says Add or Edit, so exact encoding needs confirmation.","Field baseline; codes proposed","Base FSD section 9"),
("part_no","Part Number","Required. Add must reference a Part Master item not yet registered for the ongoing period. Edit must reference an existing part; Part No cannot change.","Baseline","Base FSD BR-03 / section 9"),
("pcs_case","Pieces per case","Required for Add; may change on Edit. Positive integer per screen rule. Edit blank-cell behavior is TBD.","Field baseline; blank behavior open","Base FSD BR-04 / OP-07"),
("max_case_day","Maximum cases per day","Required for Add; may change on Edit. Positive integer. Effect on existing Mapping allocations is unresolved.","Field baseline; impact open","Base FSD BR-04 / OP-07"),
("minimum_mad","Minimum MAD","Required for Add per upload template; may change on Edit. Numeric range and blank-cell behavior are unresolved.","Field baseline; validation open","Base FSD section 9 / OP-07"),
("effective_start_date","Effective Start Date","Required for Add; may change on Edit only while part status is Candidate. Screen rule allows a minimum of today. Blank behavior on Edit is unresolved.","Field baseline; file rule open","Base FSD BR-06 / section 9 / OP-09"),
]),
("Row validation", ["ID","Validation","Rule","Status","Reference"], [
("PL-V01","Row marker","Each row has one recognized Add or Edit value. Draft codes A/E need confirmation.","Proposed detail","FSD section 9; UISS Part List"),
("PL-V02","Part number exists","Part No must exist in Part Master. For Add it must not already exist in the ongoing-period Part List; for Edit target must exist.","Baseline-derived","BR-03; UISS Part List"),
("PL-V03","Duplicate key","Part No appears at most once per workbook.","Proposed","UISS Part List"),
("PL-V04","Required Add fields","All six data fields are required for Add.","Baseline","Base FSD section 9"),
("PL-V05","Positive quantities","Pcs/Case and Max Case/Day are whole numbers greater than zero. Numeric upper limits and accepted formatting are TBD.","Constraint baseline; format/limits open","Base FSD BR-04 / OP-07"),
("PL-V06","Minimum MAD","Require a valid numeric value for Add. Whether zero/negative is allowed and maximum are TBD.","Required field baseline; rule open","Base FSD section 9 / OP-07"),
("PL-V07","Effective date","Add requires a valid date not earlier than today. Edit can change date only while Candidate.","Baseline-derived; Excel serialization open","Base FSD BR-06"),
("PL-V08","Edit status","Status is not an upload column. Edit does not change Candidate/Active/Inactive status.","Baseline-derived","Base FSD BR-07; six-column template"),
("PL-V09","Edit blanks","Draft recommendation: blank editable cells on Edit retain the current value; confirm if this applies to all editable fields.","Proposed","UISS Part List; base FSD silent"),
]),
("Processing and save behavior", ["Step","Behavior","Description","Status","Reference"], [
("1","Receive","Store workbook and create UPLOAD_ID.","Proposed","UISS Part List"),
("2","Check file","Apply common checks and verify six expected headers.","Mixed","Sections 2 and 3"),
("3","Validate rows","Run all PL-V rules; attach errors to Excel row and column.","Proposed","UISS Part List"),
("4","Reject invalid file","If any row fails, reject and write no Part List records under draft atomic-save behavior.","Proposed","SH-06"),
("5","Add part","Create ongoing-period row; look up Part Name; set Candidate and No Sync; calculate Total Pcs/Day = Pcs/Case x Max Case/Day.","Status/calculation baseline; batch path proposed","Base FSD BR-02 / BR-05"),
("6","Edit part","Update allowed values only; retain Part No, Part Name and status. Recalculate total when quantities change.","Field restrictions baseline; blank/save rules proposed","Base FSD BR-05 / BR-07"),
("7","Change metadata","For changed rows set Sync Status to No Sync and Changed By/Date to uploader/server time. Unchanged rows remain untouched.","Baseline save rule; audit fields proposed","Base FSD BR-09 / OP-10"),
("8","Finalize","Commit changes together; produce result. Database error rolls back the workbook.","Proposed","UISS Part List"),
]),
("Result counters", ["Counter","Definition","Status"], [
("TOTAL_ROWS","Data rows read from Upload worksheet.","Proposed"),
("SAVED_ROWS","Add rows inserted plus Edit rows changed.","Proposed"),
("NO_CHANGE_ROWS","Edit rows skipped because no business value changed.","Proposed"),
("ERROR_ROWS","Rows with one or more validation errors.","Proposed"),
]),
])

sheet("4 Mapping Upload", "4. Mapping per Destination Upload", "Draft for a proposed upload capability. The base FSD defines Mapping export and screen editing, but not Mapping upload.", [20,31,69,24,28], [
("Scope and proposed template", ["ID / Field","Requirement","Description","Status","Reference"], [
("MP-01 Capability","New function","Add an Upload action for D/S users in Mapping per Destination. Entry point, permission details and screen response are not in current FSD/prototype.","Proposed","Base FSD sections 1, 5, 9; OP-11"),
("MP-02 Scope","Period and parts","Apply only to ongoing period and Active parts. Period is inferred by batch, not supplied in the file.","Business rule baseline; upload proposed","Base FSD BR-12 / BR-13"),
("MP-03 Format","Workbook","Draft uses .xlsx, max 5 MB and 5,000 rows. Current Mapping export is .xls HTML and cannot be assumed to be an upload template.","Proposed","Base FSD section 9; UISS Mapping"),
("MP-04 Sheet","Worksheet","Draft worksheet name is Upload. Exact header validation and extra-sheet/column behavior need confirmation.","Proposed","UISS Mapping"),
("part_no","Part Number","Required. Must identify an Active part in ongoing period. Rows are grouped by Part No.","Proposed field; eligibility baseline","BR-12 / UISS Mapping"),
("dest_code","Destination Code","Required destination from destination master; unique within a Part No group.","Proposed field; uniqueness baseline","BR-14; master source TBD"),
("case_day","Case/Day","Required whole-number allocation. Group total must not exceed Max Case/Day; lower total is allowed.","Field proposed; sum rule baseline","BR-16"),
("total_case_day","Group total","UISS draft includes this as read-only/ignored. Recommendation: calculate from case_day and do not require an ignored column; decide template before approval.","Open","UISS Mapping; base FSD has no upload field"),
]),
("Group validation", ["ID","Validation","Rule","Status","Reference"], [
("MP-V01","Part eligibility","Part No exists for ongoing period and has Active status.","Baseline-derived","Base FSD BR-12 / BR-13"),
("MP-V02","Destination validity","Each Dest Code exists in approved destination master; physical source is unspecified.","Validation proposed; source open","Base FSD destination codes; UISS Mapping"),
("MP-V03","Unique destination","Destination code occurs once per Part No group.","Baseline","Base FSD BR-14"),
("MP-V04","Allocation value","Case/Day is a whole number. Numeric range and whether zero is valid need confirmation.","Open","Base FSD OP-07"),
("MP-V05","Group sum","Sum per Part No does not exceed Max Case/Day. A lower total is allowed and is an incomplete mapping.","Baseline","Base FSD BR-16 / BR-10"),
("MP-V06","No-destination group","Export emits blank Dest Code and Case/Day 0 for an unmapped part, but the screen blocks saving when no destination exists. Whether an upload can use blank/zero to clear a group is unresolved.","Open / conflicting rules","Base FSD section 9 / BR-16; UISS Mapping"),
("MP-V07","Part grouping","A Part No may have multiple rows; each group must satisfy row/group rules. A Part No absent from the file is unchanged.","Proposed","UISS Mapping"),
("MP-V08","Duplicate pair","Duplicate Part No/Dest Code pairs are rejected; identify both conflicting rows where feasible.","Proposed","Derived from BR-14"),
]),
("Processing and save behavior", ["Step","Behavior","Description","Status","Reference"], [
("1","Receive","Store workbook, create UPLOAD_ID and invoke or queue batch.","Proposed","UISS Mapping"),
("2","Check file","Apply proposed file controls and verify agreed Upload template.","Proposed","Sections 2 and 4"),
("3","Validate","Validate each row/group; resolve Active status, destination validity and Max Case/Day.","Proposed implementation of baseline rules","BR-12 to BR-16"),
("4","Reject invalid data","Under draft atomic-save behavior, any validation error rejects workbook and saves no destination/parent changes.","Proposed","SH-06"),
("5","Compare groups","Compare each valid destination set with ongoing-period data. Skip identical groups; absent groups stay unchanged.","Proposed","UISS Mapping"),
("6","Replace changed group","Replace all destination rows for each changed Part No in ongoing period. No-destination representation must be decided first.","Proposed / empty group open","UISS Mapping; BR-16"),
("7","Update parent","For changed groups set parent Sync Status to No Sync and update uploader/time audit fields.","Save rule baseline; audit proposed","Base FSD BR-09 / OP-10"),
("8","Finalize","Commit changed groups in one transaction and generate result. Database failure rolls back all groups.","Proposed","UISS Mapping"),
]),
("Result counters", ["Counter","Definition","Status"], [
("TOTAL_ROWS","Data rows read from Upload worksheet.","Proposed"),
("SAVED_ROWS","Changed Part No groups committed.","Proposed"),
("NO_CHANGE_ROWS","Part No groups identical to current mappings and skipped.","Proposed"),
("ERROR_ROWS","Rows with row errors; group failures identify affected rows.","Proposed"),
]),
])

sheet("5 Data and Results", "5. Data, Results and Audit", "Logical data effects and batch reporting. Physical schema and retention remain design decisions.", [25,36,72,24,27], [
("Logical data sources and effects", ["Entity","Use","Expected effect","Status","Reference"], [
("Part master","Lookup Part No/Part Name for Part List Add; eligibility for Mapping.","Read-only lookup. Authoritative physical source and outage behavior need confirmation.","Source open","Base FSD BR-03 / OP-08"),
("Part List master","Current values, status, ongoing period and Max Case/Day.","Insert valid Add; update valid Edit; recalculate total; changed rows become No Sync.","Logical entity baseline; schema TBD","UISS Part List; BR-02/05/09"),
("Mapping master","Current destinations and Case/Day per part/period.","Replace destinations for changed included groups; absent/unchanged groups untouched.","Logical entity baseline; schema TBD","UISS Mapping; BR-12 to BR-16"),
("Destination master","Validate Dest Code and obtain metadata if needed.","Read-only lookup; physical source/table not identified in base FSD.","Open","UISS Mapping"),
("Upload log","Track UPLOAD_ID, function, file, status, counts, timestamps and result location.","One logical record per workbook; proposed table TB_R_CD_UPLOAD_LOG.","Proposed","UISS Batch drafts"),
("Upload staging","Retain raw values, Excel row, validation status and errors during processing.","Separate logical staging by function; proposed names TB_T_CD_UPLOAD_PART / TB_T_CD_UPLOAD_MAPPING.","Proposed","UISS Batch drafts"),
]),
("Result workbook", ["Sheet / field","Draft content","Status","Decision required"], [
("Summary","UPLOAD_ID, function, filename, status, total/saved/unchanged/error counts and start/end time.","Proposed","Confirm counters and whether every outcome has a report."),
("Result","One line per input row (Part List) or Part No group (Mapping), with action, outcome and reason.","Proposed","Confirm exact columns and row/group layout."),
("Re-upload","Copy of input data plus error detail for correction and re-upload.","Proposed","Confirm whether all original columns are preserved."),
("Changes","Before/after values for persisted changes.","Proposed","Confirm access and data exposure requirements."),
("Availability","Make report downloadable by submitting user after completion/rejection.","Open","Confirm UI notification, expiry and access control."),
("Report failure","Finalize batch status and log report failure; show whether master changes committed.","Proposed","Confirm response when report fails after commit."),
]),
("Audit and operational controls", ["Concern","Requirement / proposal","Status","Reference"], [
("Uploader and time","Record authenticated user and server timestamps for changed data and batch log.","Proposed","Base FSD OP-10"),
("History","Trace changes to UPLOAD_ID; user-facing history granularity and old/new values undecided.","Open","Base FSD OP-03"),
("Retention","Define retention for source workbook, report, log and staging data; cleanup must respect report expiry.","Open","UISS Batch drafts"),
("Concurrent updates","Define behavior if a part/mapping changes between validation and commit. Draft recommendation: recheck current values in transaction and reject conflicts.","Proposed / open","Not covered by base FSD"),
("Duplicate submission","Define whether a repeated workbook is a new batch, content-deduplicated or blocked by request ID.","Open","Not covered by base FSD"),
("Execution","Choose synchronous vs queued processing, timeout, progress visibility and technical retry policy.","Open","UISS Batch drafts"),
]),
])

sheet("6 Messages and AC", "6. Messages and Acceptance Criteria", "Draft messages and observable outcomes. Message IDs and final copy are not approved.", [18,34,78,28], [
("Draft messages", ["ID","Context","Draft text","Status"], [
("BU-MSG-01","File rejected","File validation failed: <reason>. Check the template and upload the file again.","Proposed"),
("BU-MSG-02","Row rejected","Row <row>, column <field>: <reason>. Correct the result workbook and upload again.","Proposed"),
("BU-MSG-03","Batch completed","Upload <upload_id> completed. Total: <total>; saved: <saved>; unchanged: <unchanged>; errors: <errors>.","Proposed"),
("BU-MSG-04","Batch rejected","Upload <upload_id> was rejected. No master data was saved. Review the result workbook.","Proposed; depends on atomic-save approval"),
("BU-MSG-05","System failure","The upload could not be processed. Contact the system owner with Upload ID <upload_id>.","Proposed"),
("BU-MSG-06","Report unavailable","Processing finished, but the result report could not be created. Contact the system owner with Upload ID <upload_id>.","Proposed"),
]),
("Acceptance criteria - Part List", ["ID","Given / when","Expected result","Status"], [
("PL-AC-01","Given a valid Add row for an eligible Part No, when batch completes,","Candidate part is created for ongoing period; Part Name is resolved; Total Pcs/Day is calculated; Sync Status is No Sync.","Draft; baseline-derived"),
("PL-AC-02","Given an Add row with missing/invalid required data, when validation runs,","Report identifies Excel row/field; no workbook rows are written under proposed atomic-save rule.","Proposed"),
("PL-AC-03","Given an Edit row for an existing part with changed editable values, when batch completes,","Only allowed fields change; Part No, Part Name and status remain; total and change metadata are updated.","Baseline-derived; blank behavior open"),
("PL-AC-04","Given an Edit row with no business changes, when batch completes,","Count as unchanged; audit and Sync fields remain unchanged.","Proposed"),
("PL-AC-05","Given an Edit row changing Effective Start Date for Active/Inactive part, when validation runs,","Reject row because date can change only while Candidate.","Baseline"),
("PL-AC-06","Given one valid and one invalid row in a workbook, when batch completes,","Reject batch and change no master row if atomic-save rule is approved.","Proposed"),
]),
("Acceptance criteria - Mapping", ["ID","Given / when","Expected result","Status"], [
("MP-AC-01","Given a valid group for an Active part in ongoing period, when batch completes,","Replace destinations for included Part No if changed; absent groups remain untouched.","Proposed"),
("MP-AC-02","Given a destination repeated within the same Part No group, when validation runs,","Reject group and identify conflicting input row(s).","Baseline-derived; detail proposed"),
("MP-AC-03","Given group total above Max Case/Day, when validation runs,","Reject group.","Baseline"),
("MP-AC-04","Given group total below Max Case/Day, when batch completes,","Group may be saved and is an incomplete mapping under BR-10.","Baseline"),
("MP-AC-05","Given Candidate/Inactive part or past period, when validation runs,","Reject row/group; save no Mapping for that part.","Baseline-derived"),
("MP-AC-06","Given one invalid and one valid group in workbook, when batch completes,","Change no destination/parent row if atomic-save rule is approved.","Proposed"),
("MP-AC-07","Given a group identical to current Mapping, when batch completes,","Count unchanged; do not update parent Sync Status or audit fields.","Proposed"),
]),
])

sheet("7 Open Points", "7. Open Points and Decisions", "Resolve these items before treating this draft as a build-ready functional specification.", [14,33,76,72,14], [
("Decisions required", ["ID","Topic","Finding / impact","Decision required / draft recommendation","Impact"], [
("BU-OP-01","Approve Mapping upload","Base FSD specifies Mapping export and interactive edit, not upload.","Confirm feature, D/S ownership, entry point and access behavior.","High"),
("BU-OP-02","Mapping file compatibility","Existing Mapping export is HTML with .xls extension; proposal uses .xlsx.","Choose independent .xlsx template or define compatible export.","High"),
("BU-OP-03","Mapping template fields","UISS draft includes ignored total_case_day; base export does not specify it.","Prefer part_no, dest_code and case_day only; calculate totals in batch. Confirm.","High"),
("BU-OP-04","Clear Mapping","Export has blank Dest Code / Case-Day 0 for unmapped parts, but BR-16 blocks saving when there is no destination.","Decide whether blank/zero is a sentinel to clear a group or unmapped groups cannot be uploaded.","High"),
("BU-OP-05","Atomicity and partial success","UISS proposes all-or-nothing; base FSD does not define upload transaction behavior.","Approve all-or-nothing or define partial commit semantics. Draft recommends all-or-nothing.","High"),
("BU-OP-06","Part List marker","Base FSD says Add/Edit; UISS uses A/E.","Confirm exact template values and display wording.","Medium"),
("BU-OP-07","Edit blank cells","Base FSD does not say whether blank Edit values clear, retain or fail.","Draft recommends blanks retain current values; define if any field may be cleared.","High"),
("BU-OP-08","Numeric rules","OP-07 leaves ranges/upper bounds undefined for part quantities, MAD and Case/Day.","Specify integer/decimal, zero/negative, maximum, separators and overflow behavior.","High"),
("BU-OP-09","Effective date","Add requires a date; screen minimum is today. Excel serialization and Edit blank semantics are undefined.","Confirm accepted Excel dates and Candidate-only changes; reject dates before today.","Medium"),
("BU-OP-10","Max Case/Day and Mapping","Changing Max Case/Day can make existing allocations exceed the new maximum.","Define cross-function validation: block edit, permit inconsistent state or require Mapping correction.","High"),
("BU-OP-11","Mapping zero allocation","BR-16 and no-destination export row have conflicting implications for empty/zero allocations.","Clarify whether Case/Day must be positive and how an unmapped part is represented.","High"),
("BU-OP-12","Master data sources","Physical Part and Destination master sources and refresh behavior are unknown.","Identify source, key lengths, inactive/expired behavior and outage response.","Medium"),
("BU-OP-13","Batch invocation","Sync/queue behavior, timeout, progress and technical retry are unspecified.","Choose execution model and final status/report retrieval behavior.","Medium"),
("BU-OP-14","Result workbook","Sheet names are proposed; exact layout, generation failure and download expiry unspecified.","Approve fields, filename, storage/access, retention and report requirement.","Medium"),
("BU-OP-15","Audit and retention","History events and retention are open; staging and source/result retention also unspecified.","Define row/group/batch event granularity, old/new values, retention and cleanup.","Medium"),
("BU-OP-16","Concurrent and duplicate uploads","No rule covers simultaneous edits or duplicate workbook submission.","Define conflict detection and idempotency. Draft recommends rejecting stale conflicts.","Medium"),
("BU-OP-17","Schema and messages","UISS physical names, field lengths, message IDs and log format remain proposed/TBD.","Align logical requirements with DB convention, deployment and message catalog.","Medium"),
]),
])

wb.properties.title = "FSD - Crossdock Master Setting Batch Uploads"
wb.properties.subject = "Functional specification for Part List and Mapping per Destination batch uploads"
wb.properties.creator = "OpenAI Codex"
wb.properties.description = "Draft addendum to the existing Crossdock Master Setting FSD. Proposed behavior and open decisions are identified."
wb.active = 0
wb.save(OUT)
check = load_workbook(OUT, read_only=True, data_only=True)
print("Created:", OUT.resolve())
print("Sheets:", ", ".join(check.sheetnames))
for ws in check.worksheets:
    print(f"{ws.title}: {ws.max_row} rows x {ws.max_column} columns")
