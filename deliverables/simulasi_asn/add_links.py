"""Internal hyperlinks between related tables of simulasi_asn_rev.xlsx (same sheet or another sheet)."""
import sys, openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font
P = sys.argv[1]
wb = openpyxl.load_workbook(P)
s1, s2, st = wb['Sheet1'], wb['Sheet2'], wb['#TEMP_INSTRUCTION']

def link(cell, sheet, coord, tip):
    if cell.value in (None, 'NULL'): return
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.title}'!{coord}", tooltip=tip)
    f = cell.font
    cell.font = Font(name=f.name, size=f.size, bold=f.bold, color='FF0563C1', underline='single')

def tables(ws):
    """{table name: [(header row, {col name: col idx}, [data rows])]}"""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (c.value.startswith('TB_') or c.value.startswith('NSP.TB_')):
                hr = c.row + (2 if ws is s1 and c.value.startswith(('TB_', 'NSP.')) and ws.cell(c.row + 1, c.column).value != 'CUSTOMER_CD' and ws.cell(c.row + 1, c.column).value != 'Created' else 1)
                hr = c.row + 1
                cols = {}
                k = c.column
                while ws.cell(hr, k).value is not None:
                    cols[ws.cell(hr, k).value] = k; k += 1
                rows, r = [], hr + 1
                while ws.cell(r, c.column).value is not None:
                    rows.append(r); r += 1
                out.setdefault(c.value, []).append((hr, cols, rows))
    return out

T2 = tables(s2)
part = T2['TB_R_PRA_DAILY_ORDER_PART'][0]; mapd = T2['TB_R_PRA_DAILY_ORDER_MAPPED'][0]
ins = T2['TB_R_INSTRUCTION'][0]; his = T2['TB_H_CROSSDOCK_ORDER_CHANGE'][0]
v = lambda ws, t, r, name: ws.cell(r, t[1][name]).value
cc = lambda ws, t, r, name: ws.cell(r, t[1][name]).coordinate
po_part = {v(s2, part, r, 'PO_NO'): r for r in part[2]}
po_map = {v(s2, mapd, r, 'PO_NO'): r for r in mapd[2]}

# Sheet1 replicate blocks: date, rows
blocks = []
for row in s1.iter_rows(min_col=1, max_col=1):
    c = row[0]
    if c.value == 'NSP.TB_R_H0253_ORDER_CONTROL_PRE':
        hr = c.row + 1; cols = {s1.cell(hr, k).value: k for k in range(1, 15)}
        rows, r = [], hr + 1
        while s1.cell(r, 1).value is not None: rows.append(r); r += 1
        blocks.append((s1.cell(c.row - 1, 1).value, s1.cell(c.row - 2, 1).value, cols, rows))
def rep_row(date, co, last=True, asn=None):
    cand = [(b, r) for b in blocks if b[0] == date for r in b[3] if s1.cell(r, b[2]['CUSTOMER_ORDER_NO']).value == co
            and (asn is None or s1.cell(r, b[2]['ASN NO (dari TB_R_INSTRUCTION)']).value == asn)]
    if not cand: return None
    b, r = cand[-1] if last else cand[0]
    return f"B{r}"
created = lambda r: str(v(s2, ins, r, 'Created'))[:10]

# Sheet2 PO check table: PO_NO / SPLIT PO → PRA PART
for r in range(4, 4 + 10):
    c = s2.cell(r, 5)
    if c.value in po_part:
        link(c, s2, cc(s2, part, po_part[c.value], 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_PART')
        sp = s2.cell(r, 9)
        if isinstance(sp.value, str) and sp.value[:8] in po_part: link(sp, s2, cc(s2, part, po_part[sp.value[:8]], 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_PART (split PO)')
# Sheet2: PO_NO / MANIFEST_NO → PRA PART; PRA PART PO_NO → MAPPED; MAPPED PO_NO → PRA PART
for r in part[2]:
    link(s2.cell(r, part[1]['PO_NO']), s2, cc(s2, mapd, po_map[v(s2, part, r, 'PO_NO')], 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_MAPPED')
for r in mapd[2]:
    pr = po_part[v(s2, mapd, r, 'PO_NO')]
    link(s2.cell(r, mapd[1]['PO_NO']), s2, cc(s2, part, pr, 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_PART')
    link(s2.cell(r, mapd[1]['MANIFEST_NO']), s2, cc(s2, part, pr, 'MANIFEST_NO'), 'TB_R_PRA_DAILY_ORDER_PART')
ins_by_id = {}
for r in ins[2]:
    ins_by_id.setdefault(v(s2, ins, r, 'INSTRUCTION_ID'), []).append(r)
    pr = po_part[v(s2, ins, r, 'PO_NO')]
    link(s2.cell(r, ins[1]['PO_NO']), s2, cc(s2, part, pr, 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_PART')
    link(s2.cell(r, ins[1]['MANIFEST_NO']), s2, cc(s2, part, pr, 'MANIFEST_NO'), 'TB_R_PRA_DAILY_ORDER_PART')
    t = rep_row(created(r), v(s2, ins, r, 'CUSTOMER_ORDER_NO'), asn=v(s2, ins, r, 'ASN_NO'))
    if t: link(s2.cell(r, ins[1]['CUSTOMER_ORDER_NO']), s1, t, 'NSP.TB_R_H0253_ORDER_CONTROL_PRE (Sheet1)')
# TB_H: ASN / instruction ids → TB_R_INSTRUCTION; customer orders → Sheet1 (last block)
for r in his[2]:
    fr, to = ins_by_id[v(s2, his, r, 'FROM_INSTRUCTION_ID')][0], ins_by_id[v(s2, his, r, 'TO_INSTRUCTION_ID')][-1]
    link(s2.cell(r, his[1]['FROM_INSTRUCTION_ID']), s2, cc(s2, ins, fr, 'INSTRUCTION_ID'), 'TB_R_INSTRUCTION (before)')
    link(s2.cell(r, his[1]['TO_INSTRUCTION_ID']), s2, cc(s2, ins, to, 'INSTRUCTION_ID'), 'TB_R_INSTRUCTION (after)')
    link(s2.cell(r, his[1]['ASN_NO']), s2, cc(s2, ins, to, 'ASN_NO'), 'TB_R_INSTRUCTION')
    d = str(v(s2, his, r, 'Created'))[:10]
    for k in ('FROM_CUSTOMER_ORDER_NO', 'TO_CUSTOMER_ORDER_NO'):
        t = rep_row(d, v(s2, his, r, k))
        if t: link(s2.cell(r, his[1][k]), s1, t, 'NSP.TB_R_H0253_ORDER_CONTROL_PRE (Sheet1)')
# Sheet1: ASN NO → TB_R_INSTRUCTION row (same customer order and ASN, latest state up to that date)
dkey = lambda d: tuple(reversed(str(d).split('/')))
for date, title, cols, rows in blocks:
    for r in rows:
        asn, co = s1.cell(r, cols['ASN NO (dari TB_R_INSTRUCTION)']).value, s1.cell(r, cols['CUSTOMER_ORDER_NO']).value
        hit = [x for x in ins[2] if v(s2, ins, x, 'ASN_NO') == asn and v(s2, ins, x, 'CUSTOMER_ORDER_NO') == co and dkey(created(x)) <= dkey(date)]
        if hit: link(s1.cell(r, cols['ASN NO (dari TB_R_INSTRUCTION)']), s2, cc(s2, ins, hit[-1], 'ASN_NO'), 'TB_R_INSTRUCTION (Sheet2)')
# Sheet1: G-System block customer order → NSP replicate row of the same date (after XD08)
for row in s1.iter_rows(min_col=1, max_col=1):
    c = row[0]
    if c.value == 'TB_R_H0253_ORDER_CONTROL_PRE':
        date, r = s1.cell(c.row - 1, 1).value, c.row + 2
        while s1.cell(r, 1).value is not None:
            t = rep_row(date, s1.cell(r, 2).value, last=False)
            if t: link(s1.cell(r, 2), s1, t, 'NSP.TB_R_H0253_ORDER_CONTROL_PRE (setelah XD08)')
            r += 1
# TEMP: PO_NO / MANIFEST_NO → PRA PART; CUSTOMER_ORDER_NO → Sheet1 replicate block used by the Search
hdr = None
for r in range(1, st.max_row + 1):
    a = st.cell(r, 1).value
    if isinstance(a, str) and '[Search]' in a: date = a[:10]
    if a == 'SORT_SEQ': hdr = {st.cell(r, k).value: k for k in range(1, 18)}; continue
    if hdr and isinstance(a, int):
        po = st.cell(r, hdr['PO_NO']).value
        if po:
            link(st.cell(r, hdr['PO_NO']), s2, cc(s2, part, po_part[po], 'PO_NO'), 'TB_R_PRA_DAILY_ORDER_PART (Sheet2)')
            link(st.cell(r, hdr['MANIFEST_NO']), s2, cc(s2, part, po_part[po], 'MANIFEST_NO'), 'TB_R_PRA_DAILY_ORDER_PART (Sheet2)')
        t = rep_row(date, st.cell(r, hdr['CUSTOMER_ORDER_NO']).value, last=False)
        if t: link(st.cell(r, hdr['CUSTOMER_ORDER_NO']), s1, t, 'NSP.TB_R_H0253_ORDER_CONTROL_PRE (Sheet1)')
st['A2'] = 'Teks biru bergaris bawah = hyperlink ke tabel terkait (PO_NO / MANIFEST_NO → TB_R_PRA_DAILY_ORDER_PART, CUSTOMER_ORDER_NO → H0253 replicate di Sheet1).'
st['A2'].font = Font(italic=True, color='FF0563C1')
wb.save(P)
n = sum(1 for ws in wb for row in ws.iter_rows() for c in row if c.hyperlink)
print('links', n)
