"""Build XD01-XD15 User Requirement & Business Logic summary workbook."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from ur_data import FUNCS, DETAIL, BL_ASN, OPEN

OUT = os.path.join(os.path.dirname(__file__), '..', 'XD01-XD15_UR_Business_Logic_Summary.xlsx')
HDR = PatternFill('solid', fgColor='1F3864')
GREY = PatternFill('solid', fgColor='D9D9D9')
CAT_FILL = {'User Requirement': 'DDEBF7', 'Business Logic': 'E2EFDA', 'Validation': 'FFF2CC',
            'Abnormality / Exception': 'FCE4D6', 'Message / Log': 'EDEDED'}
THIN = Side(style='thin', color='808080')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def table(ws, row, headers, rows, widths):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, color='FFFFFF'), HDR, BOX
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[cell.column_letter].width = w
    for r in rows:
        row += 1
        for c, v in enumerate(r, 1):
            cell = ws.cell(row, c, v)
            cell.alignment, cell.border = WRAP, BOX
    return row


def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = sub
    ws['A2'].font = Font(italic=True, color='595959')


wb = Workbook()
ws = wb.active
ws.title = 'Cover'
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 90
title(ws, 'Crossdock Optimization – XD01–XD15 User Requirement & Business Logic Summary', 'GTOPAS / NSP')
info = [('Project', 'Crossdock Optimization (NSP)'), ('Document', 'Summary User Requirement & Business Logic per function'),
        ('Author', 'FID Rayhan'), ('Revision', '0.1 Initial Draft'), ('Date', '27-Oct-2026'),
        ('Source', 'UISS XD01–XD15 (0.1), function_list_v2 md, simulasi_asn_rev.xlsx (16–22 Oct), BPF Level 2'),
        ('Sheets', '1. Summary – 1 baris per function\n2. Detail UR & BL – rincian per function (filter Func ID / Category)\n'
                   '3. Business Logic ASN – rule 1–6\n4. Open Points'),
        ('Catatan', 'XD12 dihapus (digabung ke XD08 + XD11); XD14 tidak dipakai.')]
for i, (k, v) in enumerate(info, 4):
    ws.cell(i, 1, k).font = Font(bold=True)
    ws.cell(i, 1).fill = GREY
    ws.cell(i, 2, v).alignment = WRAP
    for c in (1, 2):
        ws.cell(i, c).border = BOX
r = len(info) + 5
ws.cell(r, 1, 'Legend Category').font = Font(bold=True)
for i, (k, col) in enumerate(CAT_FILL.items(), r + 1):
    ws.cell(i, 1, k).fill = PatternFill('solid', fgColor=col)
    ws.cell(i, 1).border = BOX

ws = wb.create_sheet('Summary')
title(ws, 'Summary per Function', 'User Requirement (tujuan) dan Business Logic utama setiap function')
hd = ['No', 'Func ID', 'Function Name', 'Type', 'User / PIC', 'Trigger / Schedule', 'User Requirement (Objective)',
      'Key Business Logic', 'Input', 'Output', 'Related Function', 'Status']
rows = [(i,) + f for i, f in enumerate(FUNCS, 1)]
last = table(ws, 4, hd, rows, [5, 8, 26, 11, 22, 22, 45, 60, 32, 32, 16, 14])
for rr in range(5, last + 1):
    if ws.cell(rr, 12).value in ('Removed', 'Not used'):
        for c in range(1, 13):
            ws.cell(rr, c).fill = GREY
            ws.cell(rr, c).font = Font(color='808080')
ws.freeze_panes = 'D5'
ws.auto_filter.ref = f'A4:L{last}'

ws = wb.create_sheet('Detail UR & BL')
title(ws, 'Detail User Requirement & Business Logic', 'Contoh dari simulasi_asn_rev.xlsx (PART-X, 6 pcs/case, supplier 50221-1)')
names = {f[0]: f[1] for f in FUNCS}
rows = [(i, d[0], names[d[0]]) + d[1:] for i, d in enumerate(DETAIL, 1)]
last = table(ws, 4, ['No', 'Func ID', 'Function Name', 'Category', 'Item', 'Description', 'Example (Simulation)', 'Related Table'],
             rows, [5, 8, 26, 18, 26, 75, 40, 34])
for rr in range(5, last + 1):
    ws.cell(rr, 4).fill = PatternFill('solid', fgColor=CAT_FILL[ws.cell(rr, 4).value])
ws.freeze_panes = 'E5'
ws.auto_filter.ref = f'A4:H{last}'

ws = wb.create_sheet('Business Logic ASN')
title(ws, 'Business Logic ASN (24–27 Oct)', 'Berlaku untuk XD08–XD11; contoh dari simulasi')
table(ws, 4, ['No', 'Rule', 'Penjelasan', 'Function', 'Contoh Simulasi'], BL_ASN, [5, 40, 65, 16, 50])
ws.freeze_panes = 'A5'

ws = wb.create_sheet('Open Points')
title(ws, 'Open Points', 'Keputusan yang masih menunggu')
table(ws, 4, ['ID', 'Function', 'Point', 'PIC', 'Status'], OPEN, [8, 14, 80, 18, 18])

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print('saved', os.path.abspath(OUT), len(FUNCS), 'functions', len(DETAIL), 'details')
