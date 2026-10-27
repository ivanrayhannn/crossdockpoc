"""Build the User Requirement workbook of one function (business level, FSD style).

Usage: python3 build_ur.py XD01   -> XD01_UR_<Name>.xlsx in this folder (content in ur_<id>.py)."""
import importlib, os, sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
NAVY, GREY, LIGHT = '1F3864', 'D9D9D9', 'F2F2F2'
THIN = Side(style='thin', color='808080')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
F_HEAD, F_BODY, F_BOLD = Font(name='Arial', size=9, bold=True, color='FFFFFF'), Font(name='Arial', size=9), Font(name='Arial', size=9, bold=True)
PRIO = {'Must': 'FCE4D6', 'Should': 'FFF2CC', 'Could': 'E2EFDA'}


def head(ws, title, sub, d):
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'{d.INFO["id"]} {d.INFO["name"]} – User Requirement'
    ws['A1'].font = Font(name='Arial', size=9, color='595959')
    ws['A2'] = title
    ws['A2'].font = Font(name='Arial', size=14, bold=True, color=NAVY)
    if sub:
        ws['A3'] = sub
        ws['A3'].font = Font(name='Arial', size=9, italic=True, color='595959')
    return 5


def table(ws, r, cols, widths, rows, fills=None):
    for c, (t, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(r, c, t)
        cell.font, cell.fill, cell.border = F_HEAD, PatternFill('solid', fgColor=NAVY), BOX
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if w:
            ws.column_dimensions[cell.column_letter].width = w
    for i, row in enumerate(rows):
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border, cell.alignment = (F_BOLD if c == 1 else F_BODY), BOX, WRAP
            if fills and fills[i] and fills[i][1] == c:
                cell.fill = PatternFill('solid', fgColor=fills[i][0])
        cw = [ws.column_dimensions[ws.cell(r, c + 1).column_letter].width or 10 for c in range(len(row))]
        lines = max(len(str(v)) / max(cw[c] * 1.15, 1) + str(v).count('\n') for c, v in enumerate(row))
        ws.row_dimensions[r].height = max(15, 12.5 * (int(lines) + 1))
    return r + 2


def section(ws, r, text):
    ws.cell(r, 1, text).font = Font(name='Arial', size=10, bold=True, color=NAVY)
    return r + 1


def build(fid):
    d = importlib.import_module(f'ur_{fid.lower()}')
    wb = Workbook()
    # --- Cover
    ws = wb.active
    ws.title = 'Cover'
    r = head(ws, 'User Requirement Document', 'Business-level requirement (what and why). Technical design (tables, fields, data mapping, log) is in the UISS.', d)
    info = [('Project', 'Crossdock Optimization (NSP)'), ('Function', f'{d.INFO["id"]} {d.INFO["name"]} ({d.INFO["type"]})'),
            ('Related functions', d.INFO['related']), ('Business process (BPF)', d.INFO['bpf']), ('Technical design', f'UISS {d.INFO["uiss"]}'),
            ('Author', 'FID Rayhan'), ('Revision', '0.1 Initial Draft'), ('Date', '28-Oct-2026'), ('Status', 'Draft for user review')]
    for col, w in zip('ABCDE', [24, 18, 18, 50, 22]):
        ws.column_dimensions[col].width = w
    for k, v in info:
        ws.cell(r, 1, k).font, ws.cell(r, 1).fill, ws.cell(r, 1).border = F_BOLD, PatternFill('solid', fgColor=GREY), BOX
        c = ws.cell(r, 2, v); c.font, c.alignment = F_BODY, WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        for cc in range(2, 6): ws.cell(r, cc).border = BOX
        r += 1
    r = section(ws, r + 1, 'Revision history')
    r = table(ws, r, ['Rev', 'Date', 'Author', 'Description', 'Approved by'], [None] * 5, [['0.1', '28-Oct-2026', 'FID Rayhan', 'Initial Draft', '']])
    r = section(ws, r, 'Sign-off')
    r = table(ws, r, ['Role', 'Name', 'Date', 'Signature'], [None] * 4,
              [['Procurement (user)', '', '', ''], ['D/S Team (user)', '', '', ''], ['Business System Analyst', 'FID Rayhan', '', '']])
    for rr in range(r - 4, r - 1): ws.row_dimensions[rr].height = 30
    r = section(ws, r, 'Contents')
    for i, t in enumerate(['1. Overview (purpose, scope, users, business process)', '2. Functional Requirements', '3. Business Rules', '4. Screen Overview',
                           '5. Messages', '6. Non-Functional', '7. Acceptance Criteria', '8. Open Points'], r):
        ws.cell(i, 1, t).font = F_BODY

    # --- 1 Overview
    ws = wb.create_sheet('1 Overview')
    r = head(ws, '1. Overview', '', d)
    for col, w in zip('ABCD', [16, 90, 24, 12]):
        ws.column_dimensions[col].width = w
    r = section(ws, r, '1.1 Purpose')
    ws.cell(r, 1, d.PURPOSE).alignment = WRAP; ws.cell(r, 1).font = F_BODY
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); ws.row_dimensions[r].height = 30
    r = section(ws, r + 2, '1.2 Scope')
    r = table(ws, r, ['In / Out', 'Description'], [None] * 2, [['In scope', x] for x in d.SCOPE_IN] + [['Out of scope', x] for x in d.SCOPE_OUT])
    r = section(ws, r, '1.3 Users and access')
    r = table(ws, r, ['User', 'Can do', 'Role'], [None] * 3, [[u[0], u[2], u[1]] for u in d.USERS])
    r = section(ws, r, '1.4 Business process (BPF F1)')
    table(ws, r, ['No', 'Activity', 'Who', 'BPF step'], [None] * 4, d.PROCESS)

    # --- 2 FR
    ws = wb.create_sheet('2 Functional Req')
    r = head(ws, '2. Functional Requirements', 'What the user must be able to do. Priority: Must / Should / Could.', d)
    table(ws, r, ['ID', 'Requirement', 'User', 'Priority', 'Business rules', 'BPF step'], [8, 70, 18, 9, 20, 10],
          [list(x) for x in d.FR], fills=[(PRIO[x[3]], 4) for x in d.FR])
    ws.freeze_panes = f'A{r + 1}'

    # --- 3 BR
    ws = wb.create_sheet('3 Business Rules')
    r = head(ws, '3. Business Rules', 'Rules the system must apply, in business wording.', d)
    table(ws, r, ['ID', 'Topic', 'Rule'], [8, 18, 100], [list(x) for x in d.BR])
    ws.freeze_panes = f'A{r + 1}'

    # --- 4 Screens
    ws = wb.create_sheet('4 Screen Overview')
    r = head(ws, '4. Screen Overview (conceptual)', 'Screenshots of the prototype as illustration; final layout is defined in the UISS.', d)
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 28, 110
    for cap, img, desc in d.SCREENS:
        ws.cell(r, 1, cap).font = Font(name='Arial', size=10, bold=True, color=NAVY)
        c = ws.cell(r, 2, desc); c.font, c.alignment = F_BODY, WRAP
        ws.row_dimensions[r].height = 28
        r += 2
        path = os.path.join(ROOT, img)
        if os.path.exists(path):
            pic = Image(path)
            k = 640 / pic.width
            pic.width, pic.height = int(pic.width * k), int(pic.height * k)
            ws.add_image(pic, f'B{r}')
            r += int(pic.height / 18) + 4
        else:
            ws.cell(r, 2, '[screen image]').font = F_BODY
            r += 2

    # --- 5 Messages
    ws = wb.create_sheet('5 Messages')
    r = head(ws, '5. Messages to the user', 'Business wording; message IDs and exact texts are in the UISS.', d)
    table(ws, r, ['When', 'Message', 'Rule'], [18, 90, 10], [list(x) for x in d.MESSAGES])

    # --- 6 NFR
    ws = wb.create_sheet('6 Non-Functional')
    r = head(ws, '6. Non-Functional Requirements', '', d)
    table(ws, r, ['Area', 'Requirement'], [18, 100], [list(x) for x in d.NFR])

    # --- 7 AC
    ws = wb.create_sheet('7 Acceptance Criteria')
    r = head(ws, '7. Acceptance Criteria', 'Given / When / Then; the base of the test scenarios (UAT).', d)
    table(ws, r, ['ID', 'Scenario', 'FR', 'BR', 'Result (UAT)'], [8, 90, 14, 16, 14], [list(x) + [''] for x in d.AC])
    ws.freeze_panes = f'A{r + 1}'

    # --- 8 Open
    ws = wb.create_sheet('8 Open Points')
    r = head(ws, '8. Assumptions and Open Points', '', d)
    table(ws, r, ['ID', 'Point', 'PIC', 'Status'], [8, 90, 14, 12], [list(x) for x in d.OPEN])

    for s in wb.worksheets:
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth, s.page_setup.fitToHeight = 1, 0
        s.sheet_properties.pageSetUpPr.fitToPage = True
    out = os.path.join(HERE, f'{fid}_UR_{d.INFO["name"].replace(" ", "")}.xlsx')
    wb.save(out)
    print('written', out)


if __name__ == '__main__':
    sys.path.insert(0, HERE)
    build(sys.argv[1] if len(sys.argv) > 1 else 'XD01')
