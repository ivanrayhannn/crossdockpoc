"""Build Crossdock_DB_Structure.xlsx in the layout of the DB Structure document (GTOPAS_SPS_DB_Structure): Revision, Table List and one
sheet per table (header block, Database Design / Name / Row Size / PK Row Size / Logging, Desc, field list). Columns come from the same
sources as the ERD (build_erd.load_tables). Run: python3 build_dbstructure.py (from this folder)."""
import os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from build_erd import load_tables, norm

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Crossdock_DB_Structure.xlsx')

# id, physical name, logical name, design kind, description
TABLES = [
    ('MST-01', 'TB_M_PART_CANDIDATE', 'Crossdock Part Candidate', 'Master', 'Candidate / Active / Inactive crossdock part (Part List of XD01).'),
    ('MST-02', 'TB_M_PART_CANDIDATE_DEST', 'Part Mapping per Destination', 'Master', 'Case/Day per part, period and destination (Mapping per Destination of XD01).'),
    ('HIS-01', 'TB_H_PART_CANDIDATE', 'Part Candidate History', 'History', 'Every change of a crossdock part (screen, upload, calculation, sync).'),
    ('HIS-02', 'TB_H_CROSSDOCK_ORDER_CHANGE', 'Crossdock Order Change History', 'History', 'Change Order Number (XD11): cancelled qty in an ASN moved to other lines.'),
    ('TR-01', 'TB_R_PART_DAD_RESULT', 'DAD Volume Calculation Result', 'Transaction', 'Result of XD05 process 1 per part and period.'),
    ('TR-02', 'TB_R_CROSSDOCK_IF_RESULT', 'Crossdock IF Result', 'Transaction', 'Send status of the crossdock master per part and target (dummy table).'),
    ('TR-03', 'TB_R_INSTRUCTION', 'Crossdock Instruction', 'Transaction', 'Real mapping of Create ASN: case x customer order line x PO.'),
    ('TR-04', 'TB_R_PRA_DAILY_ORDER_MAPPED', 'Pra Daily Order Mapped Qty', 'Transaction', 'Qty of a crossdock PO line mapped to customer orders.'),
    ('TMP-01', '#TEMP_INSTRUCTION', 'Crossdock Pre-mapping Temporary', 'Temporary', 'Pre-mapping of one [Search] (session).'),
    ('REP-01', 'TB_M_H0125_DESTINATION', 'Destination Master (H0125)', 'Replicated', 'Destination master replicated from DB Surrounding by XD13.'),
    ('REP-02', 'TB_R_ACTUAL_DEMAND_SUMMARY', 'Actual Demand Summary', 'Replicated', 'Actual demand replicated from DB Surrounding by XD04.'),
    ('REP-03', 'TB_R_H0253_ORDER_CONTROL_PRE', 'Order Control Info (H0253)', 'Replicated', 'Customer order lines replicated by XD08, with NSP columns HOLD_FLAG / PO_MAPPED_FLAG / UPD_STATUS (NULL / 6 / 2).'),
    ('EXT-01', 'TB_M_PART_MASTER', 'Part Master', 'Existing', 'Existing part master (reference; Crossdock uses PART_NAME).'),
    ('EXT-02', 'TB_R_PRA_DAILY_ORDER_MANIFEST', 'Pra Daily Order Manifest', 'Existing', 'Existing (reference): crossdock order = CROSS_DOCK_FLAG X.'),
    ('EXT-03', 'TB_R_PRA_DAILY_ORDER_PART', 'Pra Daily Order Part', 'Existing', 'Existing (reference): part data of the crossdock manifest.'),
    ('EXT-04', 'TB_R_DAILY_ORDER_MANIFEST', 'Daily Order Manifest', 'Existing', 'Existing (reference): ASN manifest written by Submit ASN.'),
    ('EXT-05', 'TB_R_DAILY_ORDER_PART', 'Daily Order Part', 'Existing', 'Existing (reference): ASN part x case written by Submit ASN.'),
    ('EXT-06', 'TB_R_CASE_LABEL_LAST_SEQ', 'Case Label Last Seq', 'Existing', 'Existing (reference): last sequence of the ASN number and the case label.'),
]
ID_FILL = {'MST': '92D050', 'HIS': 'B4C6E7', 'TR': 'F4B084', 'TMP': 'D9D9D9', 'REP': 'CCC0DA', 'EXT': 'BFBFBF'}
GREY = PatternFill('solid', fgColor='FFD9D9D9')
thin = Side(style='thin', color='FF000000')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F10, F10B, F8 = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True), Font(name='Arial', size=8)
SIZE = {'NUMERIC': 9, 'NUMBER': 9, 'INT': 4, 'BIGINT': 8, 'DATE': 3, 'DATETIME': 8, 'TIMESTAMP': 8}


def col(n):           # 1-based column number from a letter
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(n)


def put(ws, r, a, b, value, label=False, font=None, align='left'):
    ca, cb = col(a), col(b)
    if cb > ca:
        ws.merge_cells(start_row=r, start_column=ca, end_row=r, end_column=cb)
    for c in range(ca, cb + 1):
        ws.cell(r, c).border = BOX
        if label:
            ws.cell(r, c).fill = GREY
    cell = ws.cell(r, ca, value)
    cell.font = font or F10
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=False, shrink_to_fit=True)


def parse_type(t):
    t = t.replace(' (inferred)', '').replace('+', '')
    m = re.match(r'^([A-Z#]+)\s*\(([^,)]*)(?:,\s*([^)]*))?\)', t)
    if m:
        ln = m.group(2)
        return m.group(1), int(ln) if ln.isdigit() else ln, m.group(3) or ''
    return t, '', ''


def size_of(typ, ln):
    if typ in ('VARCHAR', 'CHAR', 'NVARCHAR'):
        return ln if isinstance(ln, int) else None
    return SIZE.get(typ, None)


def logical(name, desc):
    if ';' in desc and len(desc.split(';')[0]) <= 40:
        return desc.split(';')[0].strip()
    return name.replace('_', ' ').title().replace(' Cd', ' Code').replace(' Dt', ' Date')


def header(ws, title, page):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 0.57
    for c in range(2, col('CW') + 1):
        ws.column_dimensions[L(c)].width = 1.43
    put(ws, 2, 'B', 'G', 'Phase', True)
    put(ws, 2, 'H', 'N', 'FSD', font=F10B)
    put(ws, 2, 'O', 'AA', 'Project Name', True)
    put(ws, 2, 'AB', 'CJ', 'Crossdock Optimization (NSP)')
    put(ws, 2, 'CK', 'CN', 'Page', True)
    put(ws, 2, 'CO', 'CV', page)
    for r, lab, val, lab2, d in ((3, 'Block Name', 'Crossdock', 'Created', '20-Oct-26'), (4, 'Document Name', 'DB Structure', 'Updated', '20-Oct-26')):
        put(ws, r, 'B', 'N', lab, True)
        put(ws, r, 'O', 'AH', val)
        put(ws, r, 'AI', 'AS', 'Function Name' if r == 3 else 'Title', True)
        put(ws, r, 'AT', 'BV', 'Crossdock (XD01 – XD15)' if r == 3 else title)
        put(ws, r, 'BW', 'CB', lab2, True)
        put(ws, r, 'CC', 'CJ', d)
        put(ws, r, 'CK', 'CN', 'By', True)
        put(ws, r, 'CO', 'CV', 'FID Rayhan')


def build():
    tables, used = load_tables()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Revision'
    header(ws, 'Revision', '1')
    put(ws, 6, 'B', 'G', 'Rev', True)
    put(ws, 6, 'H', 'S', 'Date', True)
    put(ws, 6, 'T', 'CJ', 'Description', True)
    put(ws, 6, 'CK', 'CV', 'By', True)
    put(ws, 7, 'B', 'G', '0.1', align='center')
    put(ws, 7, 'H', 'S', '20-Oct-2026')
    put(ws, 7, 'T', 'CJ', 'Initial Draft: Crossdock tables (new, replicated, existing reference)')
    put(ws, 7, 'CK', 'CV', 'FID Rayhan')

    tl = wb.create_sheet('Table List')
    header(tl, 'Table List', '2')
    heads = [('B', 'F', 'No'), ('G', 'V', 'Logical Table Name'), ('W', 'AN', 'Physical Table Name'), ('AO', 'AU', 'Row Size'), ('AV', 'BB', 'PK Row Size'),
             ('BC', 'BG', 'Logging'), ('BH', 'BL', 'Max Record'), ('BM', 'BR', 'Table Size'), ('BS', 'CV', 'Description')]
    for a, b, text in heads:
        put(tl, 6, a, b, text, True, align='center')
    for n, (tid, name, lname, kind, desc) in enumerate(TABLES, 1):
        cols = tables[name][1]
        sizes, pk_sizes, unknown = 0, 0, False
        for c in cols:
            typ, ln, _ = parse_type(c[1])
            s = size_of(typ, ln)
            if s is None:
                unknown = True
                continue
            sizes += s
            if c[2] == 'Y':
                pk_sizes += s
        r = 6 + n
        tl.cell(r, 1, n).font = F8
        put(tl, r, 'B', 'F', tid, font=F8)
        for c in range(col('B'), col('F') + 1):
            tl.cell(r, c).fill = PatternFill('solid', fgColor='FF' + ID_FILL[tid.split('-')[0]])
        put(tl, r, 'G', 'V', lname, font=F8)
        put(tl, r, 'W', 'AN', name, font=F8)
        put(tl, r, 'AO', 'AU', f'{sizes}{"*" if unknown else ""}', font=F8, align='right')
        put(tl, r, 'AV', 'BB', pk_sizes, font=F8, align='right')
        put(tl, r, 'BC', 'BG', 'Yes', font=F8, align='center')
        put(tl, r, 'BH', 'BL', '', font=F8)
        put(tl, r, 'BM', 'BR', '', font=F8)
        put(tl, r, 'BS', 'CV', f'{kind}: {desc}', font=F8)
    r = 8 + len(TABLES)
    tl.cell(r, 2, '* Row Size without the columns whose length is still TBD.  MST / HIS / TR / TMP = new Crossdock table · REP = replicated from DB Surrounding · '
                  'EXT = existing NSP table (reference only, structure from GTOPAS_SPS_DB_Structure / data export)').font = Font(name='Arial', size=8, italic=True)

    for n, (tid, name, lname, kind, desc) in enumerate(TABLES, 3):
        ws = wb.create_sheet(tid)
        header(ws, lname, str(n))
        cols = tables[name][1]
        pk_order, sizes, pk_sizes = {}, 0, 0
        for c in cols:
            if c[2] == 'Y':
                pk_order[c[0]] = len(pk_order) + 1
            typ, ln, _ = parse_type(c[1])
            s = size_of(typ, ln) or 0
            sizes += s
            pk_sizes += s if c[2] == 'Y' else 0
        put(ws, 6, 'B', 'N', 'Database Design', True)
        put(ws, 6, 'O', 'AH', kind)
        put(ws, 6, 'AI', 'AN', 'Name', True)
        put(ws, 6, 'AO', 'BF', name)
        put(ws, 6, 'BG', 'BO', 'Row Size (B)', True)
        put(ws, 6, 'BP', 'BT', sizes, align='right')
        put(ws, 6, 'BU', 'CE', 'PK Row Size (B)', True)
        put(ws, 6, 'CF', 'CJ', pk_sizes, align='right')
        put(ws, 6, 'CK', 'CQ', 'Logging', True)
        put(ws, 6, 'CR', 'CV', 'Yes', align='center')
        put(ws, 7, 'B', 'G', 'Desc', True)
        put(ws, 7, 'H', 'CV', f'{lname}: {desc} Used by: {", ".join(sorted(set(used.get(name, [])))) or "-"}.')
        heads = [('B', 'D', 'No.'), ('E', 'U', 'Field Name'), ('V', 'AL', 'Logical Name'), ('AM', 'AR', 'Type'), ('AS', 'AV', 'Len'), ('AW', 'AZ', 'Prec'),
                 ('BA', 'BC', 'PK'), ('BD', 'BF', 'N-N'), ('BG', 'BM', 'Default'), ('BN', 'CV', 'Description')]
        for a, b, text in heads:
            put(ws, 9, a, b, text, True, align='center')
        for i, c in enumerate(cols, 1):
            r = 9 + i
            typ, ln, prec = parse_type(c[1])
            desc_c = c[4].split(';', 1)[1].strip() if ';' in c[4] and len(c[4].split(';')[0]) <= 40 else c[4]
            values = [('B', 'D', f'{i:02d}', 'center'), ('E', 'U', c[0], 'left'), ('V', 'AL', logical(norm(c[0]), c[4]), 'left'), ('AM', 'AR', typ, 'left'),
                      ('AS', 'AV', ln, 'center'), ('AW', 'AZ', prec, 'center'), ('BA', 'BC', pk_order.get(c[0], ''), 'center'),
                      ('BD', 'BF', 'x' if c[3] == 'N' else '', 'center'), ('BG', 'BM', '', 'left'), ('BN', 'CV', desc_c, 'left')]
            for a, b, v, al in values:
                put(ws, r, a, b, v, font=F8, align=al)
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    for w in (wb['Revision'], tl):
        w.page_setup.orientation = 'landscape'
        w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT)
    print('written', OUT)


if __name__ == '__main__':
    build()
