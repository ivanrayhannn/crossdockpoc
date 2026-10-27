"""F2 (27-Oct): the Day D+1 band goes to column AG, so header (rows 1-3), title (row 5) and legend (rows 24-25)
are widened from Y to AG; row 6 is centred. Shapes of the flow sheets are kept; run build_bpf.py xlsx afterwards (it writes the other sheets again)."""
import copy, os, sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter as L, column_index_from_string as C
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from build_bpf import add_drawings, kept_drawings

XLSX = os.path.join(os.path.dirname(HERE), 'Crossdock_BPF_Level2.xlsx')


def widen(ws, rows, old, new):
    o, n = C(old), C(new)
    for r in rows:
        for m in list(ws.merged_cells.ranges):
            if m.min_row <= r <= m.max_row and m.max_col == o:
                rng = (m.min_col, m.min_row, m.max_row)
                ws.unmerge_cells(str(m))
                ws.merge_cells(start_row=rng[1], start_column=rng[0], end_row=rng[2], end_column=n)
        src = ws.cell(r, o)
        fill, font, b = copy.copy(src.fill), copy.copy(src.font), copy.copy(src.border)
        for c in range(o, n + 1):
            cell = ws.cell(r, c)
            cell.fill, cell.font = copy.copy(fill), copy.copy(font)
            cell.border = Border(top=copy.copy(b.top), bottom=copy.copy(b.bottom), right=copy.copy(b.right) if c == n else None)


def main():
    wb = load_workbook(XLSX)
    titles = wb.sheetnames
    drawings = kept_drawings(XLSX, [t for t in titles if t.startswith('F')])  # other sheets: build_bpf.py xlsx writes them again
    ws = wb['F2']
    if ws['Z1'].border.top.style is None:
        widen(ws, [1, 2, 3, 5, 24, 25], 'Y', 'AG')
    for name in ('F1', 'F2'):
        a = wb[name]['A6'].alignment
        wb[name]['A6'].alignment = Alignment(horizontal='center', vertical=a.vertical or 'center')
    wb.save(XLSX)
    add_drawings(XLSX, titles, drawings)
    print('written', XLSX)


if __name__ == '__main__':
    main()
