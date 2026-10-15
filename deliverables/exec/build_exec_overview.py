"""Build the executive overview workbook of Crossdock Optimization and the task / confirmation md.

python3 deliverables/exec/build_exec_overview.py
  -> deliverables/exec/Crossdock_Executive_Overview.xlsx
  -> deliverables/md/00_TaskList_dan_Konfirmasi.md
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from project_items import AS_OF, CONFIRMATIONS, DONE, FUNCTIONS, TASKS  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'deliverables', 'exec', 'Crossdock_Executive_Overview.xlsx')
MD = os.path.join(ROOT, 'deliverables', 'md', '00_TaskList_dan_Konfirmasi.md')

FONT = 'Arial'
NAVY, SLATE, LIGHT, PALE, WHITE = '1F3864', '44546A', 'DDEBF7', 'F2F6FB', 'FFFFFF'
YELLOW, GREY, GREY_TXT, RED = 'FFF2CC', 'EDEDED', '7F7F7F', 'C00000'
ARROW = '5B6B82'
thin = Side(style='thin', color='BFC9D6')
med = Side(style='medium', color=NAVY)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(size=10, bold=False, color='262626', italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def title(ws, text, sub, width_cols):
    ws['B2'] = text
    ws['B2'].font = f(18, True, NAVY)
    ws['B3'] = sub
    ws['B3'].font = f(10, color=SLATE)
    for c in range(2, width_cols + 1):
        ws.cell(4, c).border = Border(bottom=Side(style='medium', color=RED))
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2


def header_row(ws, row, cols, start=2):
    for i, h in enumerate(cols):
        c = ws.cell(row, start + i, h)
        c.font = f(10, True, WHITE)
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def body(c, wrap=True, center=False, bold=False, color='262626'):
    c.font = f(10, bold, color)
    c.alignment = Alignment(horizontal='center' if center else 'left', vertical='top', wrap_text=wrap)
    c.border = BOX


def landscape(ws, area, one_page=False):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = area


wb = Workbook()

# ---------------------------------------------------------------- Function Map
fm = wb.active
fm.title = 'Function Map'
title(fm, 'Function Map', f'13 fungsi Crossdock Optimization (UISS X1-X13) · status per {AS_OF}', 12)
cols = ['ID', 'Function', 'Type', 'User / Actor', 'Phase', 'Trigger / Frequency', 'Input', 'Output', 'Prototype', 'UISS', 'MD']
header_row(fm, 6, cols)
for w, col in zip([7, 32, 9, 22, 17, 22, 30, 34, 13, 11, 11], range(2, 13)):
    fm.column_dimensions[L(col)].width = w
FM_FIRST = 7
for i, fn in enumerate(FUNCTIONS):
    r = FM_FIRST + i
    vals = list(fn) + ['Draft 0.1', 'Updated']
    for j, v in enumerate(vals):
        c = fm.cell(r, 2 + j, v)
        body(c, center=j in (0, 2, 8, 9, 10), bold=j == 0)
    fm.cell(r, 2).font = f(10, True, NAVY)
    fm.row_dimensions[r].height = 30
    if i % 2:
        for j in range(len(vals)):
            fm.cell(r, 2 + j).fill = fill(PALE)
FM_LAST = FM_FIRST + len(FUNCTIONS) - 1
fm.freeze_panes = 'C7'
landscape(fm, f'A1:L{FM_LAST + 1}')

# ---------------------------------------------------------------- Task List
tl = wb.create_sheet('Task List')
title(tl, 'Task List – Belum Dikerjakan', f'Pekerjaan yang masih terbuka per {AS_OF}. Kolom Status bisa diubah (pilihan list).', 9)
for w, col in zip([5, 58, 14, 12, 20, 22, 20, 10], range(2, 10)):
    tl.column_dimensions[L(col)].width = w
STATUSES = ['Not started', 'Waiting confirmation', 'In progress', 'Done']
tl['B6'] = 'Ringkasan status'
tl['B6'].font = f(10, True, NAVY)
TL_HDR = 12
TL_FIRST = TL_HDR + 1
TL_LAST = TL_FIRST + len(TASKS) - 1
for k, s in enumerate(STATUSES):
    lab = tl.cell(7 + k, 3, s)
    lab.font = f(10)
    lab.border = BOX
    cnt = tl.cell(7 + k, 4, f'=COUNTIF($H${TL_FIRST}:$H${TL_LAST},C{7 + k})')
    cnt.font = f(10, True, NAVY)
    cnt.alignment = Alignment(horizontal='center')
    cnt.border = BOX
tl.row_dimensions[TL_HDR - 1].height = 6
header_row(tl, TL_HDR, ['No', 'Task', 'Function', 'Category', 'PIC', 'Depends on', 'Status', 'Priority'])
dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=False)
tl.add_data_validation(dv)
for i, t in enumerate(TASKS):
    r = TL_FIRST + i
    body(tl.cell(r, 2, i + 1), center=True)
    for j, v in enumerate(t):
        body(tl.cell(r, 3 + j, v), center=j in (5, 6))
    tl.row_dimensions[r].height = 30
    tl.cell(r, 8).fill = fill(YELLOW)
    dv.add(tl.cell(r, 8))
tl.conditional_formatting.add(f'I{TL_FIRST}:I{TL_LAST}', CellIsRule(operator='equal', formula=['"High"'], font=Font(name=FONT, bold=True, color=RED)))
tl.freeze_panes = f'C{TL_FIRST}'
landscape(tl, f'A1:I{TL_LAST + 1}')

# ---------------------------------------------------------------- Konfirmasi
kf = wb.create_sheet('Konfirmasi')
title(kf, 'Daftar Konfirmasi dari User', f'Pertanyaan yang perlu keputusan user per {AS_OF}.', 11)
kf['B5'] = 'Cara isi: tulis keputusan di kolom kuning "Keputusan" dan "Tanggal" (contoh: "Setuju usulan" · 15-Oct-2026). Kolom Status berubah otomatis menjadi Decided.'
kf['B5'].font = f(9, italic=True, color=SLATE)
for w, col in zip([5, 11, 17, 46, 34, 30, 28, 13, 11], range(2, 11)):
    kf.column_dimensions[L(col)].width = w
KF_HDR = 7
header_row(kf, KF_HDR, ['No', 'Function', 'Topik', 'Pertanyaan', 'Usulan (BSA)', 'Dampak jika belum diputuskan', 'Keputusan', 'Tanggal', 'Status'])
KF_FIRST = KF_HDR + 1
KF_LAST = KF_FIRST + len(CONFIRMATIONS) - 1
for i, q in enumerate(CONFIRMATIONS):
    r = KF_FIRST + i
    body(kf.cell(r, 2, i + 1), center=True)
    for j, v in enumerate(q):
        body(kf.cell(r, 3 + j, v), center=j == 0, bold=j == 1)
    for c in (8, 9):
        body(kf.cell(r, c, None))
        kf.cell(r, c).fill = fill(YELLOW)
    kf.row_dimensions[r].height = 42
    st = kf.cell(r, 10, f'=IF(H{r}="","Open","Decided")')
    body(st, center=True, bold=True)
kf.conditional_formatting.add(f'J{KF_FIRST}:J{KF_LAST}', CellIsRule(operator='equal', formula=['"Open"'], font=Font(name=FONT, bold=True, color=RED)))
kf.conditional_formatting.add(f'J{KF_FIRST}:J{KF_LAST}', CellIsRule(operator='equal', formula=['"Decided"'], font=Font(name=FONT, bold=True, color='2E7D32')))
kf.freeze_panes = f'C{KF_FIRST}'
landscape(kf, f'A1:K{KF_LAST + 1}')

# ---------------------------------------------------------------- Business Flow (swimlane)
bf = wb.create_sheet('Business Flow')
LANES = [
    ('G-TOPAS /\nDB Surrounding', 'E7E6E6'),
    ('NSP Batch\n(System)', 'D9E1F2'),
    ('Procurement', 'E2EFDA'),
    ('D/S Team', 'FCE4D6'),
    ('Supplier', 'FFF2CC'),
    ('Global System\n(TMMIN)', 'EDE7F6'),
]
LANE_IDX = {n.split('\n')[0].split(' /')[0].split(' (')[0]: i for i, (n, _) in enumerate(LANES)}
# column -> list of (lane key, ids, text, next_scope)
G, N, P, D, S, GS = 'G-TOPAS', 'NSP Batch', 'Procurement', 'D/S Team', 'Supplier', 'Global System'
COLS = [
    [(G, 'Source', 'Actual demand &\ndestination master', False)],
    [(N, 'X5 · X11', 'Replicate data\nto NSP', False), (P, 'X1 · X2', 'Register candidate part\n(Min MAD, Pcs/Case,\nMax Case/Day)', False)],
    [(N, 'X6', 'DAD volume calc:\navg MAD ≥ Min MAD\n→ Active / Inactive', False)],
    [(N, 'X7', 'Mapping destination calc:\nsplit Max Case/Day\nper destination', False)],
    [(D, 'X3 · X4', 'Review & adjust\nmapping per\ndestination', False)],
    [(GS, 'X3 Sync (RPA)', 'Mapping synced\nto global system', False)],
    None,  # phase separator
    [(G, 'Source', 'Customer order\n(crossdock, status 7)', False)],
    [(N, 'X9 · X10', 'Replicate orders +\nsupplier PO\n(Pra Daily Order)', False)],
    [(S, 'X8 → X12', 'Get Order: order ↔ PO\npre-mapping (FIFO)\n→ Ready / Not Ready', False), (D, 'X8 detail', 'Monitor detail &\nhold order line', False)],
    [(S, 'X8', 'Create ASN (mapping)\n→ Case List\n→ Submit (Created)', False)],
    [(S, 'X13', 'ASN Inquiry &\ncase label', False)],
    [(S, 'Next scope', 'Deliver cases to\ncrossdock', True)],
]
# edges between consecutive columns: (from col idx, from lane, to lane)
EDGES = {
    0: [(G, N)],
    1: [(N, N), (P, N)],
    2: [(N, N)],
    3: [(N, D)],
    4: [(D, GS)],
    7: [(G, N)],
    8: [(N, S), (N, D)],
    9: [(S, S), (D, S)],
    10: [(S, S)],
    11: [(S, S)],
}
title(bf, 'Crossdock Optimization – Business Flow (Executive View)',
      'Alur end-to-end: A. Master setup bulanan (part, kalkulasi, mapping) → B. Operasi harian (order, PO, ASN). Kode X = UISS / fungsi.', 3 + len(COLS) * 3)
bf.column_dimensions['B'].width = 15
LANE_H, ROW0 = 5, 8
BOX_W, C1_W, C2_W = 21, 1.3, 2.4


def step_col(k):
    return 3 + k * 3


def mid_row(lane):
    return ROW0 + lane * LANE_H + 2


last_col = step_col(len(COLS) - 1)
# phase band
bf.row_dimensions[6].height = 22
for name, a, b, color in (('A. MASTER SETUP (bulanan)', 0, 5, NAVY), ('B. OPERASI HARIAN', 7, len(COLS) - 1, SLATE)):
    c0, c1 = step_col(a), step_col(b)
    bf.merge_cells(start_row=6, start_column=c0, end_row=6, end_column=c1)
    c = bf.cell(6, c0, name)
    c.font = f(11, True, WHITE)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal='center', vertical='center')
# lanes
for li, (lname, lcol) in enumerate(LANES):
    r0 = ROW0 + li * LANE_H
    bf.merge_cells(start_row=r0, start_column=2, end_row=r0 + LANE_H - 1, end_column=2)
    lc = bf.cell(r0, 2, lname)
    lc.font = f(10, True, NAVY)
    lc.fill = fill(lcol)
    lc.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for rr in range(r0, r0 + LANE_H):
        bf.row_dimensions[rr].height = 17 if rr not in (r0, r0 + LANE_H - 1) else 9
        for cc in range(3, last_col + 1):
            bf.cell(rr, cc).fill = fill('FAFBFD' if li % 2 == 0 else WHITE)
        bf.cell(rr, 2).border = Border(left=med, right=med, top=med if rr == r0 else None, bottom=med if rr == r0 + LANE_H - 1 else None)
    for cc in range(3, last_col + 1):
        bf.cell(r0 + LANE_H - 1, cc).border = Border(bottom=Side(style='thin', color='BFC9D6'))
for k in range(len(COLS)):
    bf.column_dimensions[L(step_col(k))].width = BOX_W
    bf.column_dimensions[L(step_col(k) + 1)].width = C1_W
    bf.column_dimensions[L(step_col(k) + 2)].width = C2_W
# boxes
lane_color = {LANES[i][0].split('\n')[0].split(' /')[0]: LANES[i][1] for i in range(len(LANES))}
for k, boxes in enumerate(COLS):
    if boxes is None:
        sep = step_col(k)
        bf.column_dimensions[L(sep)].width = 4
        for rr in range(ROW0, ROW0 + LANE_H * len(LANES)):
            bf.cell(rr, sep).border = Border(left=Side(style='mediumDashed', color=RED))
            bf.cell(rr, sep).fill = fill(WHITE)
        continue
    for lane, ids, text, nxt in boxes:
        li = LANE_IDX[lane]
        r0 = ROW0 + li * LANE_H + 1
        c = step_col(k)
        bf.merge_cells(start_row=r0, start_column=c, end_row=r0 + 2, end_column=c)
        cell = bf.cell(r0, c)
        idc = GREY_TXT if nxt else RED
        cell.value = CellRichText(TextBlock(InlineFont(rFont=FONT, sz=9, b=True, color=idc), ids + '\n'),
                                  TextBlock(InlineFont(rFont=FONT, sz=9, color=GREY_TXT if nxt else '262626', i=nxt), text))
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        bc = GREY if nxt else (lane_color.get(lane) or WHITE)
        side = Side(style='dashed' if nxt else 'medium', color=GREY_TXT if nxt else NAVY)
        for rr in range(r0, r0 + 3):
            bf.cell(rr, c).fill = fill(bc)
            bf.cell(rr, c).border = Border(left=side, right=side, top=side if rr == r0 else None, bottom=side if rr == r0 + 2 else None)
# connectors (elbow): vertical bar in C1 column, arrow head in C2 column at target lane
for k, edges in EDGES.items():
    c1, c2 = step_col(k) + 1, step_col(k) + 2
    rows = [mid_row(LANE_IDX[a]) for a, _ in edges] + [mid_row(LANE_IDX[b]) for _, b in edges]
    for rr in range(min(rows), max(rows) + 1):
        bf.cell(rr, c1).fill = fill(ARROW)
    for _, b in edges:
        h = bf.cell(mid_row(LANE_IDX[b]), c2, '►')
        h.font = Font(name=FONT, size=11, bold=True, color=ARROW)
        h.alignment = Alignment(horizontal='left', vertical='center')
        h.fill = fill(WHITE) if LANE_IDX[b] % 2 else fill('FAFBFD')
# phase link note under the separator
end_row = ROW0 + LANE_H * len(LANES)
note_r = end_row + 1
bf.merge_cells(start_row=note_r, start_column=step_col(5), end_row=note_r, end_column=step_col(8))
n = bf.cell(note_r, step_col(5), 'Mapping hasil Fase A dipakai untuk order crossdock (status 7) di Fase B.')
n.font = f(9, italic=True, color=RED)
n.alignment = Alignment(horizontal='center')
# key rules
kr = note_r + 2
rules_a = ['Part didaftarkan Procurement dengan Min MAD sebagai acuan.',
           'DAD calc: rata-rata MAD N bulan ≥ Min MAD → Active, selain itu Inactive.',
           'Mapping calc: Max Case/Day dibagi per destinasi sesuai porsi MAD (min. 1 case).',
           'D/S team bisa menyesuaikan mapping; perubahan dikirim ke global system via RPA.']
rules_b = ['Unit ASN = Part Number + Customer Order Date, selalu masuk 1 ASN utuh.',
           'Get Order: order dipasangkan ke PO secara FIFO (pre-mapping); mapping final saat Create ASN.',
           '1 case per Part + Order Date + destinasi; beda part = beda case.',
           'Supplier melihat level case; detail & Hold Flag order line hanya untuk D/S team.']
for (hdr, rules, c0, c1) in (('Aturan kunci – Fase A', rules_a, step_col(0), step_col(5)), ('Aturan kunci – Fase B', rules_b, step_col(7), step_col(12))):
    bf.merge_cells(start_row=kr, start_column=c0, end_row=kr, end_column=c1)
    h = bf.cell(kr, c0, hdr)
    h.font = f(10, True, WHITE)
    h.fill = fill(NAVY)
    h.alignment = Alignment(vertical='center', indent=1)
    for i, t in enumerate(rules):
        rr = kr + 1 + i
        bf.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c1)
        cc = bf.cell(rr, c0, f'{i + 1}.  {t}')
        cc.font = f(9.5)
        cc.fill = fill(PALE)
        cc.alignment = Alignment(vertical='center', indent=1)
        bf.row_dimensions[rr].height = 17
# legend
lg = kr + 6
bf.cell(lg, 2, 'Legenda').font = f(10, True, NAVY)
items = [('D9E1F2', 'Kotak = langkah / fungsi; kode merah = UISS (X1-X13)'), (ARROW, 'Panah = urutan alur'),
         (GREY, 'Abu-abu putus-putus = scope berikutnya (belum didesain)'), (WHITE, 'Garis merah putus-putus = batas fase')]
for i, (col, txt) in enumerate(items):
    c = bf.cell(lg + 1 + i, 2)
    c.fill = fill(col)
    c.border = BOX
    t = bf.cell(lg + 1 + i, 3, txt)
    t.font = f(9)
landscape(bf, f'A1:{L(last_col + 1)}{lg + 5}', one_page=True)

# ---------------------------------------------------------------- Executive Summary
ex = wb.create_sheet('Executive Summary', 0)
title(ex, 'Crossdock Optimization – Executive Summary', f'NSP · disusun oleh FID Rayhan (BSA) · status per {AS_OF}', 13)
for col in range(2, 14):
    ex.column_dimensions[L(col)].width = 13
ex['B6'] = 'Tujuan'
ex['B6'].font = f(11, True, NAVY)
ex.merge_cells('B7:M8')
ex['B7'] = ('Mengoptimalkan pengiriman part crossdock: menentukan part yang layak crossdock berdasarkan demand aktual, '
            'membagi kapasitas (Max Case/Day) ke destinasi ekspor, lalu memasangkan order customer ke PO supplier '
            'sehingga supplier bisa membuat ASN per case secara langsung.')
ex['B7'].font = f(10)
ex['B7'].alignment = Alignment(wrap_text=True, vertical='top')
ex.row_dimensions[7].height = 22
ex.row_dimensions[8].height = 22
# KPI tiles (formulas)
FMr = f"'Function Map'!$D${FM_FIRST}:$D${FM_LAST}"
tiles = [
    ('Total fungsi', f"=COUNTA('Function Map'!$B${FM_FIRST}:$B${FM_LAST})"),
    ('Screen', f'=COUNTIF({FMr},"Screen")'),
    ('Batch', f'=COUNTIF({FMr},"Batch")'),
    ('Task terbuka', f"=COUNTIF('Task List'!$H${TL_FIRST}:$H${TL_LAST},\"<>Done\")"),
    ('Konfirmasi open', f"=COUNTIF(Konfirmasi!$J${KF_FIRST}:$J${KF_LAST},\"Open\")"),
    ('Prototype screen', f"=COUNTIFS({FMr},\"Screen\",'Function Map'!$J${FM_FIRST}:$J${FM_LAST},\"Yes\")"),
]
for i, (lab, frm) in enumerate(tiles):
    c0 = 2 + i * 2
    ex.merge_cells(start_row=10, start_column=c0, end_row=10, end_column=c0 + 1)
    ex.merge_cells(start_row=11, start_column=c0, end_row=12, end_column=c0 + 1)
    a = ex.cell(10, c0, lab)
    a.font = f(9, True, WHITE)
    a.fill = fill(NAVY if i < 3 else (RED if i in (3, 4) else SLATE))
    a.alignment = Alignment(horizontal='center', vertical='center')
    v = ex.cell(11, c0, frm)
    v.font = f(22, True, NAVY if i < 3 else (RED if i in (3, 4) else SLATE))
    v.alignment = Alignment(horizontal='center', vertical='center')
    for rr in (10, 11, 12):
        for cc in (c0, c0 + 1):
            ex.cell(rr, cc).border = BOX
            if rr > 10:
                ex.cell(rr, cc).fill = fill(PALE)
ex.row_dimensions[11].height = 22
ex.row_dimensions[12].height = 22
# scope by phase
r = 14
ex.cell(r, 2, 'Lingkup per fase').font = f(11, True, NAVY)
header_row(ex, r + 1, ['Fase', '', 'Fungsi', '', '', '', 'Hasil bisnis', '', '', '', '', ''])
ex.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=3)
ex.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=7)
ex.merge_cells(start_row=r + 1, start_column=8, end_row=r + 1, end_column=13)
phases = [
    ('A. Master setup (bulanan)', 'X1, X2 Part List · X5, X11 replikasi data · X6 DAD calc · X7 mapping calc · X3, X4 mapping & sync',
     'Daftar part crossdock Active dan kapasitas case/day per destinasi untuk periode baru, tersinkron ke global system.'),
    ('B. Operasi harian', 'X9, X10 replikasi order · X8 ASN Crossdock Creation (Get Order → X12 pre-mapping, Create ASN, Submit) · X13 ASN Inquiry',
     'Supplier membuat ASN per case dari order yang sudah dapat PO; D/S team memantau detail destinasi & pcs.'),
]
for i, (p, fx, out) in enumerate(phases):
    rr = r + 2 + i
    for a, b, v in ((2, 3, p), (4, 7, fx), (8, 13, out)):
        ex.merge_cells(start_row=rr, start_column=a, end_row=rr, end_column=b)
        c = ex.cell(rr, a, v)
        body(c, bold=a == 2)
        for cc in range(a, b + 1):
            ex.cell(rr, cc).border = BOX
    ex.row_dimensions[rr].height = 42
# done & next
r = 20
ex.cell(r, 2, 'Sudah selesai').font = f(11, True, NAVY)
ex.cell(r, 8, 'Langkah berikutnya (prioritas High)').font = f(11, True, NAVY)
for i, t in enumerate(DONE):
    ex.merge_cells(start_row=r + 1 + i, start_column=2, end_row=r + 1 + i, end_column=7)
    c = ex.cell(r + 1 + i, 2, f'✔  {t}')
    c.font = f(9.5)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ex.row_dimensions[r + 1 + i].height = 30
high = [t for t in TASKS if t[6] == 'High']
for i, t in enumerate(high):
    ex.merge_cells(start_row=r + 1 + i, start_column=8, end_row=r + 1 + i, end_column=13)
    c = ex.cell(r + 1 + i, 8, f'▸  {t[0]} ({t[3]})')
    c.font = f(9.5)
    c.alignment = Alignment(wrap_text=True, vertical='top')
rr = r + 2 + max(len(DONE), len(high))
ex.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=13)
ex.cell(rr, 2, 'Detail: sheet Business Flow (alur), Function Map (13 fungsi), Task List, Konfirmasi (keputusan yang dibutuhkan). '
               'Prototype: https://ivanrayhannn.github.io/crossdockpoc/').font = f(9, italic=True, color=SLATE)
landscape(ex, f'A1:N{rr + 1}', one_page=True)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {'Executive Summary': RED, 'Business Flow': NAVY}.get(ws.title, SLATE)
wb.active = 0
wb.save(OUT)
print('saved', OUT)

# ---------------------------------------------------------------- markdown
md = [f'# Task List & Daftar Konfirmasi – Crossdock Optimization', '',
      f'Status per **{AS_OF}**. Sumber data: `deliverables/exec/project_items.py` (sama dengan Excel `deliverables/exec/Crossdock_Executive_Overview.xlsx`).', '',
      '## Index dokumen', '', '| ID | Fungsi | Tipe | User | Fase | UISS | MD |', '|---|---|---|---|---|---|---|']
md_files = sorted(os.listdir(os.path.dirname(MD)))
uiss_files = sorted(os.listdir(os.path.join(ROOT, 'deliverables', 'uiss')))
for fn in FUNCTIONS:
    u = next((x for x in uiss_files if x.startswith(fn[0] + '_')), '-')
    m = ', '.join(x for x in md_files if x.startswith(fn[0] + '_') or (fn[0] in ('X5', 'X6', 'X7') and x.startswith('X5-X7')))
    md.append(f'| {fn[0]} | {fn[1]} | {fn[2]} | {fn[3]} | {fn[4]} | `{u}` | {m or "-"} |')
md += ['', '## Sudah selesai', ''] + [f'- {t}' for t in DONE]
md += ['', '## Task list (belum dikerjakan)', '', '| No | Task | Function | Kategori | PIC | Tergantung | Status | Prioritas |', '|---|---|---|---|---|---|---|---|']
md += [f'| {i + 1} | {t[0]} | {t[1]} | {t[2]} | {t[3]} | {t[4]} | {t[5]} | {t[6]} |' for i, t in enumerate(TASKS)]
md += ['', '## Perlu konfirmasi dari user', '', '| No | Function | Topik | Pertanyaan | Usulan BSA | Dampak jika belum diputuskan |', '|---|---|---|---|---|---|']
md += [f'| {i + 1} | {q[0]} | {q[1]} | {q[2]} | {q[3]} | {q[4]} |' for i, q in enumerate(CONFIRMATIONS)]
md += ['']
open(MD, 'w').write('\n'.join(md))
print('saved', MD)
