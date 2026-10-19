"""Build a UISS workbook with the uiss-generator skill, plus two project additions:

- spec key "sample_sheet": an extra sheet (e.g. "Sample Input") with one value per column, added after the
  template sheets: {"title": "...", "intro": ["..."], "sections": [{"title": "...", "columns": ["..."],
  "widths": [..], "rows": [[...], ...], "spans": [..] (optional: sheet columns taken by each column, merged)}]}
- the template's broken defined names and external links are removed afterwards (clean_uiss_xlsx.py).

Usage: uiss_build.py SPEC.json -o OUT.xlsx
"""
import argparse, glob, json, os, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = glob.glob('/root/.claude/skills/synced/*/uiss-generator/scripts')
if not SKILL:
    raise SystemExit('uiss-generator skill not found')
sys.path.insert(0, SKILL[0])
import build_uiss  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_BODY, F_BOLD = Font(name='Arial', size=9), Font(name='Arial', size=9, bold=True)
F_TITLE = Font(name='Arial', size=12, bold=True)
HEAD = PatternFill('solid', fgColor='FFFFCC99')
GROUP = PatternFill('solid', fgColor='FFF2F2F2')


def add_sample_sheet(wb, cfg):
    ws = wb.create_sheet(cfg.get('title', 'Sample Input'))
    ws.sheet_view.showGridLines = False
    ws['A1'] = cfg.get('heading', cfg.get('title', 'Sample Input'))
    ws['A1'].font = F_TITLE
    r = 2
    for line in cfg.get('intro', []):
        ws.cell(r, 1, line).font = F_BODY
        r += 1
    widths, merged = {}, []
    for sec in cfg['sections']:
        r += 1
        ws.cell(r, 1, sec['title']).font = F_BOLD
        r += 1
        spans = sec.get('spans') or [1] * len(sec['columns'])
        starts = [1 + sum(spans[:i]) for i in range(len(spans))]

        def put(r, i, v, font, fill=None, align=None):
            c0, c1 = starts[i], starts[i] + spans[i] - 1
            for c in range(c0, c1 + 1):
                ws.cell(r, c).border = BOX
                if fill:
                    ws.cell(r, c).fill = fill
            cell = ws.cell(r, c0, v)
            cell.font, cell.alignment = font, align
            if c1 > c0:
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
                if isinstance(v, str):
                    merged.append((r, c0, c1, v))
        for i, title in enumerate(sec['columns']):
            put(r, i, title, F_BOLD, HEAD, Alignment(horizontal='center', vertical='center', wrap_text=True))
        if not sec.get('spans'):
            for c, w in enumerate(sec.get('widths', []), 1):
                widths[c] = max(widths.get(c, 0), w)
        r += 1
        for row in sec['rows']:
            shade = isinstance(row, dict) and row.get('group')
            values = row['values'] if isinstance(row, dict) else row
            for i in range(len(sec['columns'])):
                v = values[i] if i < len(values) else None
                put(r, i, v if v != '' else None, F_BODY, GROUP if shade else None,
                    Alignment(vertical='top', wrap_text=True, horizontal='right' if isinstance(v, (int, float)) else 'left'))
            r += 1
    for c, w in widths.items():
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    # Excel does not fit the height of a merged cell: estimate the number of lines
    heights = {}
    for r, c0, c1, v in merged:
        chars = sum(widths.get(c, 9) for c in range(c0, c1 + 1)) * 1.15
        lines = sum(max(1, -(-len(p) // max(int(chars), 1))) for p in v.split('\n'))
        heights[r] = max(heights.get(r, 1), lines)
    for r, n in heights.items():
        if n > 1:
            ws.row_dimensions[r].height = 12.5 * n
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('spec')
    ap.add_argument('-o', '--out', required=True)
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding='utf-8'))
    sample = spec.get('sample_sheet')
    if sample:
        orig = build_uiss.Book.finalize

        def finalize(self, out_path):
            real_save = self.wb.save

            def save(path):
                add_sample_sheet(self.wb, sample)
                real_save(path)
            self.wb.save = save
            orig(self, out_path)
        build_uiss.Book.finalize = finalize
    res = build_uiss.build(spec, a.out)
    subprocess.run([sys.executable, os.path.join(HERE, 'clean_uiss_xlsx.py'), a.out], check=True, stdout=subprocess.DEVNULL)
    print(json.dumps({k: res[k] for k in res if k != 'sheets'} | {'sheets': res.get('sheets', []) + ([sample.get('title', 'Sample Input')] if sample else [])}, indent=1))


if __name__ == '__main__':
    main()
